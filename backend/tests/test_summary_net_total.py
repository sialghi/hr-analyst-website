import pandas as pd
from openpyxl import Workbook

from app.pipeline.excel_writer import FMT_RUPIAH, _sheet_summary_overview
from app.pipeline.service import _compute_summary_overview_json


def test_summary_json_calculates_net_total_and_places_it_last():
    name = "Test Person"
    df_prep = pd.DataFrame([{
        "Cabang": "Jakarta",
        "Nama": name,
        "Profil": "OFFICE",
        "Jam_Masuk": None,
        "Jam_Keluar": None,
        "Status_Data": "Lengkap",
        "Kategori_Hari": "Kerja",
    }])
    df_telat = pd.DataFrame(columns=["Nama", "Durasi_Telat_Jam"])
    df_lembur = pd.DataFrame([{
        "Nama": name,
        "Jam_Lembur_Bulat": 1,
        "Bonus_Lembur_Rp": 200,
    }])
    df_pulang_duluan = pd.DataFrame(columns=["Nama", "Durasi_Kurang_Jam"])
    df_tidak_masuk = pd.DataFrame(columns=["Nama", "Jml_Hari_Tidak_Masuk"])
    df_uang_makan = pd.DataFrame([{
        "Nama": name,
        "Uang_Makan_Base": 1200,
        "Potongan_Telat": 0,
        "Potongan_Pulang_Duluan": 0,
        "Bonus_Tanggal_Merah": 0,
        "Total_Uang_Makan": 1000,
    }])
    adjustments = {name.lower(): (30, 25)}
    bpjs = {"test person": {"bpjs_kesehatan": 50, "bpjs_tk": 100}}

    result, _ = _compute_summary_overview_json(
        df_prep, df_telat, df_lembur, df_pulang_duluan, df_tidak_masuk,
        df_uang_makan, adjustments_dict=adjustments, bpjs_dict=bpjs,
    )

    assert result["headers"][-1] == "Total Bersih (Rp)"
    assert result["types"]["Total Bersih (Rp)"] == "currency"
    assert result["rows"][0]["Total Bersih (Rp)"] == 1055


def test_summary_workbook_adds_net_total_formula_as_last_column():
    workbook = Workbook()
    employees = pd.DataFrame([{
        "Cabang": "Jakarta",
        "Nama": "Test Person",
        "Profil": "OFFICE",
    }])

    _sheet_summary_overview(
        workbook,
        employees,
        adjustments_dict={"test person": (30, 25)},
        bpjs_dict={"test person": {"bpjs_kesehatan": 50, "bpjs_tk": 100}},
    )

    worksheet = workbook["2. Summary Overview"]
    assert worksheet.cell(row=4, column=31).value == "Total Bersih (Rp)"
    assert worksheet.cell(row=5, column=31).value == "=AD5-U5-V5-X5+W5+Y5"
    assert worksheet.cell(row=5, column=31).number_format == FMT_RUPIAH
    assert worksheet.column_dimensions["AE"].width == 22
