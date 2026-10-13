# -*- coding: utf-8 -*-
import datetime
import io
import asyncio
import unittest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base, get_db
from app.main import app
from app.models import (
    User, Employee, EmploymentContract, EmployeeEmploymentTransition, Profile, RoleEnum,
    LeaveRequest, LeaveBalanceLedger, EmployeeImportCandidate,
)
from app.config import CONTRACT_REMINDER_DAYS, calculate_tenure, get_today_jakarta
from app.auth import get_current_user, require_hr_master
from app.routers import whatsapp

from sqlalchemy.pool import StaticPool
from openpyxl import Workbook, load_workbook

SQLALCHEMY_DATABASE_URL = "sqlite:///:memory:"
engine = create_engine(SQLALCHEMY_DATABASE_URL, connect_args={"check_same_thread": False}, poolclass=StaticPool)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)


def override_get_db():
    try:
        db = TestingSessionLocal()
        yield db
    finally:
        db.close()


def override_auth():
    return User(id=1, nama="HR Master Tester", email="hr@test.com", role=RoleEnum.hr_master)


app.dependency_overrides[get_db] = override_get_db
app.dependency_overrides[get_current_user] = override_auth
app.dependency_overrides[require_hr_master] = override_auth

client = TestClient(app)


class TestContractManagement(unittest.TestCase):

    def setUp(self):
        Base.metadata.drop_all(bind=engine)
        Base.metadata.create_all(bind=engine)
        db = TestingSessionLocal()
        p = Profile(code="OFFICE", nama="Office Staf", hari_kerja=["Senin", "Selasa", "Rabu", "Kamis", "Jumat"])
        db.add(p)
        db.commit()
        db.close()

    def test_1_calculate_tenure_logic(self):
        today = datetime.date(2026, 10, 1)

        # Joined 2 years 5 months ago (2024-05-01)
        t1 = calculate_tenure(datetime.date(2024, 5, 1), target_date=today)
        self.assertEqual(t1["years"], 2)
        self.assertEqual(t1["months"], 5)
        self.assertEqual(t1["display"], "2 tahun 5 bulan")

        # Joined 6 months ago (2026-04-01)
        t2 = calculate_tenure(datetime.date(2026, 4, 1), target_date=today)
        self.assertEqual(t2["years"], 0)
        self.assertEqual(t2["months"], 6)
        self.assertEqual(t2["display"], "6 bulan")

        # Empty join_date
        t3 = calculate_tenure(None)
        self.assertEqual(t3["display"], "-")

    def test_2_contract_date_validation_error(self):
        payload = {
            "nama": "Budi Kontrak Invalid",
            "profile_code": "OFFICE",
            "employment_status": "PKWT",
            "join_date": "2026-01-01",
            "contract_start_date": "2026-10-01",
            "contract_end_date": "2026-09-01",  # Invalid: end_date sebelum start_date
        }
        response = client.post("/employees", json=payload)
        self.assertEqual(response.status_code, 400)
        self.assertIn("setelah tanggal mulai", response.json()["detail"])

    def test_3_expiring_contracts_h10_reminder(self):
        today = get_today_jakarta()

        # Karyawan 1: Kontrak berakhir 5 hari lagi (H-5) -> EXPIRING_SOON
        t1_start = today - datetime.timedelta(days=100)
        t1_end = today + datetime.timedelta(days=5)

        res1 = client.post("/employees", json={
            "nama": "Siti PKWT H-5",
            "profile_code": "OFFICE",
            "employment_status": "PKWT",
            "join_date": "2025-01-01",
            "contract_start_date": t1_start.isoformat(),
            "contract_end_date": t1_end.isoformat(),
        })
        self.assertEqual(res1.status_code, 200)

        # Karyawan 2: Kontrak berakhir 30 hari lagi -> ACTIVE (Aman)
        t2_start = today
        t2_end = today + datetime.timedelta(days=30)
        client.post("/employees", json={
            "nama": "Andi PKWT Aman",
            "profile_code": "OFFICE",
            "employment_status": "PKWT",
            "join_date": "2026-01-01",
            "contract_start_date": t2_start.isoformat(),
            "contract_end_date": t2_end.isoformat(),
        })

        # Query endpoint expiring contracts H-10
        exp_res = client.get("/employees/contracts/expiring")
        self.assertEqual(exp_res.status_code, 200)
        data = exp_res.json()

        self.assertEqual(data["reminder_days_config"], 10)
        self.assertEqual(data["total"], 1)
        self.assertEqual(data["data"][0]["nama"], "Siti PKWT H-5")
        self.assertEqual(data["data"][0]["days_left"], 5)
        self.assertEqual(data["data"][0]["status_label"], "EXPIRING_SOON")

    def test_4_pkwt_promotion_to_permanent(self):
        today = get_today_jakarta()
        t_start = today - datetime.timedelta(days=180)
        t_end = today + datetime.timedelta(days=180)

        # Create PKWT
        res = client.post("/employees", json={
            "nama": "Reno Promosi",
            "profile_code": "OFFICE",
            "employment_status": "PKWT",
            "join_date": "2026-01-01",
            "contract_start_date": t_start.isoformat(),
            "contract_end_date": t_end.isoformat(),
        })
        self.assertEqual(res.status_code, 200)
        emp_id = res.json()["id"]

        # Promote to PKWTT
        update_res = client.put(f"/employees/{emp_id}", json={
            "nama": "Reno Promosi",
            "profile_code": "OFFICE",
            "employment_status": "PKWTT",
            "join_date": "2026-01-01",
        })
        self.assertEqual(update_res.status_code, 200)
        updated_emp = update_res.json()

        self.assertEqual(updated_emp["employment_status"], "PKWTT")
        # Kontrak lama PKWT tersimpan di riwayat dengan status PROMOTED_TO_PERMANENT
        self.assertEqual(len(updated_emp["contracts"]), 1)
        self.assertEqual(updated_emp["contracts"][0]["status"], "PROMOTED_TO_PERMANENT")
        self.assertIn("Diangkat menjadi PKWTT", updated_emp["contracts"][0]["keterangan"])

    def test_phl_to_pkwt_records_phl_period_and_active_contract(self):
        db = TestingSessionLocal()
        employee = Employee(
            employee_code="EMP-PHL-TO-PKWT",
            nama="Rina PHL ke PKWT",
            profile_code="OFFICE",
            employment_status="PHL",
            join_date=datetime.date(2022, 3, 1),
            active=True,
        )
        db.add(employee)
        db.commit()
        employee_id = employee.id
        db.close()

        response = client.put(f"/employees/{employee_id}", json={
            "nama": "Rina PHL ke PKWT",
            "profile_code": "OFFICE",
            "employment_status": "PKWT",
            "join_date": "2022-03-01",
            "phl_period_start_date": "2022-03-15",
            "contract_start_date": "2024-07-01",
            "contract_end_date": "2025-07-01",
            "contract_keterangan": "PKWT pertama",
        })

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["employment_status"], "PKWT")
        self.assertEqual(data["join_date"], "2022-03-01")
        self.assertEqual(data["contracts"][0]["contract_number"], 1)
        self.assertEqual(data["contracts"][0]["status"], "ACTIVE")
        self.assertIn("Perubahan status PHL ke PKWT", data["contracts"][0]["keterangan"])
        self.assertEqual(len(data["phl_to_pkwt_transitions"]), 1)
        self.assertEqual(data["phl_to_pkwt_transitions"][0]["phl_start_date"], "2022-03-15")
        self.assertEqual(data["phl_to_pkwt_transitions"][0]["phl_end_date"], "2024-06-30")
        self.assertEqual(data["phl_to_pkwt_transitions"][0]["transition_date"], "2024-07-01")
        self.assertEqual(
            data["phl_to_pkwt_transitions"][0]["employment_contract_id"],
            data["contracts"][0]["id"],
        )

        db = TestingSessionLocal()
        saved = db.query(Employee).filter(Employee.id == employee_id).one()
        self.assertEqual(saved.join_date, datetime.date(2022, 3, 1))
        self.assertEqual(len(saved.contracts), 1)
        self.assertEqual(len(saved.phl_to_pkwt_transitions), 1)
        db.close()

    def create_phl_transition_employee(self, employee_code="EMP-HISTORY-EDIT"):
        employee = Employee(
            employee_code=employee_code,
            nama="Karyawan Histori",
            profile_code="OFFICE",
            employment_status="PHL",
            join_date=datetime.date(2022, 3, 1),
            active=True,
        )
        db = TestingSessionLocal()
        db.add(employee)
        db.commit()
        employee_id = employee.id
        db.close()

        response = client.put(f"/employees/{employee_id}", json={
            "nama": "Karyawan Histori",
            "profile_code": "OFFICE",
            "employment_status": "PKWT",
            "join_date": "2022-03-01",
            "phl_period_start_date": "2022-03-15",
            "contract_start_date": "2024-07-01",
            "contract_end_date": "2025-07-01",
            "contract_keterangan": "PKWT pertama",
        })
        self.assertEqual(response.status_code, 200)
        data = response.json()
        return employee_id, data["contracts"][0]["id"], data["phl_to_pkwt_transitions"][0]["id"]

    def test_edit_contract_synchronizes_linked_phl_dates_and_validates_ranges(self):
        employee_id, contract_id, transition_id = self.create_phl_transition_employee()

        response = client.put(
            f"/employees/{employee_id}/contracts/{contract_id}",
            json={
                "start_date": "2024-08-01",
                "end_date": "2025-08-01",
                "status": "ACTIVE",
                "keterangan": "Tanggal diperbaiki",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["contract_number"], 1)
        self.assertEqual(response.json()["start_date"], "2024-08-01")

        db = TestingSessionLocal()
        transition = db.query(EmployeeEmploymentTransition).filter_by(id=transition_id).one()
        self.assertEqual(transition.transition_date, datetime.date(2024, 8, 1))
        self.assertEqual(transition.phl_end_date, datetime.date(2024, 7, 31))
        db.close()

        invalid_range = client.put(
            f"/employees/{employee_id}/contracts/{contract_id}",
            json={
                "start_date": "2024-08-01",
                "end_date": "2024-08-01",
                "status": "ACTIVE",
                "keterangan": None,
            },
        )
        invalid_status = client.put(
            f"/employees/{employee_id}/contracts/{contract_id}",
            json={
                "start_date": "2024-08-01",
                "end_date": "2025-08-01",
                "status": "INVALID",
                "keterangan": None,
            },
        )
        self.assertEqual(invalid_range.status_code, 400)
        self.assertEqual(invalid_status.status_code, 400)

    def test_edit_phl_period_updates_start_and_rejects_start_after_contract(self):
        employee_id, contract_id, transition_id = self.create_phl_transition_employee()

        updated = client.put(
            f"/employees/{employee_id}/phl-transitions/{transition_id}",
            json={"phl_start_date": "2022-04-01"},
        )
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.json()["phl_start_date"], "2022-04-01")
        self.assertEqual(updated.json()["phl_end_date"], "2024-06-30")
        self.assertEqual(updated.json()["transition_date"], "2024-07-01")

        invalid = client.put(
            f"/employees/{employee_id}/phl-transitions/{transition_id}",
            json={"phl_start_date": "2024-07-01"},
        )
        self.assertEqual(invalid.status_code, 400)

        foreign_employee = Employee(
            employee_code="EMP-HISTORY-OTHER",
            nama="Karyawan Lain",
            profile_code="OFFICE",
        )
        db = TestingSessionLocal()
        db.add(foreign_employee)
        db.commit()
        foreign_employee_id = foreign_employee.id
        db.close()
        not_owned = client.put(
            f"/employees/{foreign_employee_id}/phl-transitions/{transition_id}",
            json={"phl_start_date": "2022-04-01"},
        )
        self.assertEqual(not_owned.status_code, 404)

    def test_delete_phl_record_keeps_contract_and_delete_linked_active_contract_reverts_status(self):
        employee_id, contract_id, transition_id = self.create_phl_transition_employee()

        deleted_transition = client.delete(
            f"/employees/{employee_id}/phl-transitions/{transition_id}",
        )
        self.assertEqual(deleted_transition.status_code, 200)
        db = TestingSessionLocal()
        employee = db.query(Employee).filter_by(id=employee_id).one()
        self.assertEqual(employee.employment_status, "PKWT")
        self.assertEqual(db.query(EmploymentContract).filter_by(id=contract_id).count(), 1)
        self.assertEqual(db.query(EmployeeEmploymentTransition).filter_by(id=transition_id).count(), 0)
        db.close()

        employee_id, contract_id, transition_id = self.create_phl_transition_employee(
            "EMP-HISTORY-DELETE",
        )
        deleted_contract = client.delete(
            f"/employees/{employee_id}/contracts/{contract_id}",
        )
        self.assertEqual(deleted_contract.status_code, 200)
        db = TestingSessionLocal()
        employee = db.query(Employee).filter_by(id=employee_id).one()
        self.assertEqual(employee.employment_status, "PHL")
        self.assertEqual(db.query(EmploymentContract).filter_by(id=contract_id).count(), 0)
        self.assertEqual(db.query(EmployeeEmploymentTransition).filter_by(id=transition_id).count(), 0)
        db.close()

    def test_invalid_phl_to_pkwt_transition_does_not_change_employee_or_create_history(self):
        db = TestingSessionLocal()
        employee = Employee(
            employee_code="EMP-PHL-INVALID",
            nama="Budi PHL",
            profile_code="OFFICE",
            employment_status="PHL",
            join_date=datetime.date(2023, 1, 1),
            active=True,
        )
        db.add(employee)
        db.commit()
        employee_id = employee.id
        db.close()

        base_payload = {
            "nama": "Budi PHL",
            "profile_code": "OFFICE",
            "employment_status": "PKWT",
            "join_date": "2023-01-01",
            "phl_period_start_date": "2023-02-01",
        }
        missing_dates = client.put(f"/employees/{employee_id}", json=base_payload)
        self.assertEqual(missing_dates.status_code, 400)

        overlapping_dates = client.put(
            f"/employees/{employee_id}",
            json={
                **base_payload,
                "contract_start_date": "2023-02-01",
                "contract_end_date": "2024-02-01",
            },
        )
        self.assertEqual(overlapping_dates.status_code, 400)

        db = TestingSessionLocal()
        unchanged = db.query(Employee).filter(Employee.id == employee_id).one()
        self.assertEqual(unchanged.employment_status, "PHL")
        self.assertEqual(unchanged.join_date, datetime.date(2023, 1, 1))
        self.assertEqual(db.query(EmploymentContract).filter_by(employee_id=employee_id).count(), 0)
        self.assertEqual(db.query(whatsapp.models.EmployeeEmploymentTransition).filter_by(employee_id=employee_id).count(), 0)
        db.close()

    def test_5_contract_renewal(self):
        today = get_today_jakarta()
        t1_start = today - datetime.timedelta(days=100)
        t1_end = today + datetime.timedelta(days=10)

        res = client.post("/employees", json={
            "nama": "Dewi Perpanjangan",
            "profile_code": "OFFICE",
            "employment_status": "PKWT",
            "join_date": "2025-10-01",
            "contract_start_date": t1_start.isoformat(),
            "contract_end_date": t1_end.isoformat(),
        })
        emp_id = res.json()["id"]

        # Tambahkan perpanjangan kontrak baru (#2)
        t2_start = t1_end + datetime.timedelta(days=1)
        t2_end = t2_start + datetime.timedelta(days=180)

        renew_res = client.post(f"/employees/{emp_id}/contracts", json={
            "start_date": t2_start.isoformat(),
            "end_date": t2_end.isoformat(),
            "keterangan": "Perpanjangan Kontrak Ke-2 (6 Bulan)",
        })
        self.assertEqual(renew_res.status_code, 200)
        new_c = renew_res.json()
        self.assertEqual(new_c["contract_number"], 2)
        self.assertEqual(new_c["status"], "ACTIVE")

        # Fetch employee detail and verify history
        detail_res = client.get(f"/employees/{emp_id}")
        emp_data = detail_res.json()
        self.assertEqual(len(emp_data["contracts"]), 2)
        # Contract #2 ACTIVE, Contract #1 RENEWED
        statuses = {c["contract_number"]: c["status"] for c in emp_data["contracts"]}
        self.assertEqual(statuses[1], "RENEWED")
        self.assertEqual(statuses[2], "ACTIVE")

    def test_6_master_workbook_dry_run_and_apply(self):
        employee = Employee(
            employee_code="EMP-MASTER-001",
            nama="Rina Cabang",
            profile_code="OFFICE",
            employment_status="PKWTT",
            active=True,
        )
        db = TestingSessionLocal()
        db.add(employee)
        db.commit()
        db.close()

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Pembaruan Master"
        sheet.append([
            "Kode Karyawan",
            "Nama Karyawan",
            "NIK",
            "Cabang",
            "Kode Profil",
            "Status Kepegawaian",
            "Tanggal Masuk",
            "Status Aktif",
        ])
        sheet.append([
            "EMP-MASTER-001",
            "Rina Cabang",
            "3201010101010001",
            "Cab. Bandung",
            "OFFICE",
            "PKWTT",
            datetime.date(2024, 1, 15),
            "Aktif",
        ])
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)

        dry_run = client.post(
            "/employees/import/nik?dry_run=true",
            files={"file": ("master.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(dry_run.status_code, 200)
        self.assertTrue(dry_run.json()["dry_run"])
        self.assertEqual(dry_run.json()["total_updated"], 1)

        db = TestingSessionLocal()
        unchanged = db.query(Employee).filter(Employee.employee_code == "EMP-MASTER-001").one()
        self.assertIsNone(unchanged.cabang)
        self.assertIsNone(unchanged.nik)
        self.assertEqual(unchanged.employment_status, "PKWTT")
        db.close()

        output.seek(0)
        applied = client.post(
            "/employees/import/nik",
            files={"file": ("master.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(applied.status_code, 200)
        self.assertFalse(applied.json()["dry_run"])

        db = TestingSessionLocal()
        updated = db.query(Employee).filter(Employee.employee_code == "EMP-MASTER-001").one()
        self.assertEqual(updated.cabang, "Cab. Bandung")
        self.assertEqual(updated.nik, "3201010101010001")
        self.assertEqual(updated.join_date, datetime.date(2024, 1, 15))
        self.assertEqual(updated.employment_status, "PKWTT")
        db.close()

    def test_7_legacy_payroll_import_template_is_rejected(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "HO"
        sheet.append(["Status", "Nama", "Uang Makan ", "Gaji Pokok "])
        sheet.append(["Office A", "Pegawai Lama", 0, None])
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)

        response = client.post(
            "/employees/import",
            files={"file": ("data-uangmakan-posisi.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("Format workbook tidak dikenali", response.json()["detail"])

    def test_8_whatsapp_login_links_phone_to_nik_after_confirmation(self):
        db = TestingSessionLocal()
        employee = Employee(
            nama="Karyawan WhatsApp",
            nik="3201010101010001",
            profile_code="OFFICE",
            active=True,
        )
        db.add(employee)
        db.commit()
        db.refresh(employee)

        sent = []

        async def fake_send(to_number, message):
            sent.append((to_number, message))
            return {"mocked": True}

        async def fake_send_list(**kwargs):
            sent.append((kwargs["to_number"], kwargs))
            return {"mocked": True}

        original_send = whatsapp.send_whatsapp_text
        original_send_list = whatsapp.send_whatsapp_list
        whatsapp.send_whatsapp_text = fake_send
        whatsapp.send_whatsapp_list = fake_send_list
        whatsapp.user_states.pop("628123456789", None)
        try:
            asyncio.run(whatsapp.process_whatsapp_incoming("08123456789", "login", None, db))
            self.assertEqual(whatsapp.user_states["08123456789"]["step"], "login_waiting_nik")

            asyncio.run(whatsapp.process_whatsapp_incoming("08123456789", "3201010101010001", None, db))
            self.assertEqual(whatsapp.user_states["08123456789"]["step"], "login_confirmation")

            asyncio.run(whatsapp.process_whatsapp_incoming("08123456789", "YaKiN Dan SaDaR", None, db))
            identity = db.query(whatsapp.models.WhatsAppIdentity).one()
            self.assertEqual(identity.phone_number, "628123456789")
            self.assertEqual(identity.employee_id, employee.id)
            self.assertEqual(identity.status, "ACTIVE")
            login_menu = next(message for _, message in sent if isinstance(message, dict))
            guide_message = next(
                message
                for _, message in sent
                if isinstance(message, str) and "Panduan HR WhatsApp" in message
            )
            self.assertTrue(guide_message.startswith("*Panduan HR WhatsApp*"))
            self.assertIn("Cuti atau Izin", guide_message)
            self.assertIn("Absen Luar", guide_message)
            self.assertIn("Status pengajuan", guide_message)
            self.assertIn("Profil", guide_message)
            self.assertIn("Slip gaji per periode", guide_message)
            self.assertIn("Periksa kembali tanggal dan keterangan", guide_message)
            self.assertIn("Ketik *panduan*", guide_message)
            self.assertIs(sent[-2][1], guide_message)
            self.assertIs(sent[-1][1], login_menu)
            self.assertEqual(
                [row["id"] for row in login_menu["sections"][0]["rows"]],
                ["menu_cuti", "menu_izin", "menu_wfl", "menu_status", "menu_profile", "menu_slip_gaji"],
            )
            self.assertEqual(
                [row["title"] for row in login_menu["sections"][0]["rows"]],
                ["Cuti", "Izin", "Absen Luar", "Status", "Profil", "Slip Gaji"],
            )
            self.assertIn("Login berhasil", login_menu["body_text"])

            sent.clear()
            asyncio.run(whatsapp.process_whatsapp_incoming("08123456789", "login", None, db))
            self.assertIn("Panduan HR WhatsApp", sent[-2][1])
            self.assertIsInstance(sent[-1][1], dict)
        finally:
            whatsapp.send_whatsapp_text = original_send
            whatsapp.send_whatsapp_list = original_send_list
            whatsapp.user_states.pop("08123456789", None)
            db.close()

    def test_whatsapp_login_rebinds_hr_revoked_phone_to_new_employee_nik(self):
        db = TestingSessionLocal()
        previous_employee = Employee(
            nama="Karyawan Sebelumnya",
            nik="3201010101010081",
            profile_code="OFFICE",
            active=True,
        )
        new_employee = Employee(
            nama="Karyawan Pengganti",
            nik="3201010101010082",
            profile_code="OFFICE",
            active=True,
        )
        db.add_all([previous_employee, new_employee])
        db.flush()
        identity = whatsapp.models.WhatsAppIdentity(
            phone_number="628123456789",
            employee_id=previous_employee.id,
            status="REVOKED",
        )
        db.add(identity)
        db.commit()
        identity_id = identity.id
        previous_employee_id = previous_employee.id
        new_employee_id = new_employee.id
        sent = []

        async def fake_send(to_number, message):
            sent.append((to_number, message))
            return {"mocked": True}

        async def fake_send_list(**kwargs):
            sent.append((kwargs["to_number"], kwargs))
            return {"mocked": True}

        original_send = whatsapp.send_whatsapp_text
        original_send_list = whatsapp.send_whatsapp_list
        whatsapp.send_whatsapp_text = fake_send
        whatsapp.send_whatsapp_list = fake_send_list
        whatsapp.user_states.pop("08123456789", None)
        try:
            asyncio.run(whatsapp.process_whatsapp_incoming(
                "08123456789", "login", None, db
            ))
            asyncio.run(whatsapp.process_whatsapp_incoming(
                "08123456789", "3201010101010082", None, db
            ))
            self.assertEqual(
                whatsapp.user_states["08123456789"]["step"],
                "login_confirmation",
            )
            asyncio.run(whatsapp.process_whatsapp_incoming(
                "08123456789", "yakin dan sadar", None, db
            ))

            db.refresh(identity)
            self.assertEqual(identity.id, identity_id)
            self.assertNotEqual(previous_employee_id, new_employee_id)
            self.assertEqual(identity.employee_id, new_employee_id)
            self.assertEqual(identity.status, "ACTIVE")
            success_audit = db.query(whatsapp.models.WhatsAppAuditLog).filter_by(
                phone_number="628123456789",
                event_type="LOGIN_SUCCESS",
            ).one()
            self.assertEqual(success_audit.employee_id, new_employee_id)
            self.assertTrue(any(
                isinstance(message, str) and "Panduan HR WhatsApp" in message
                for _, message in sent
            ))
            self.assertTrue(any(
                isinstance(message, dict) and "Login berhasil" in message["body_text"]
                for _, message in sent
            ))
        finally:
            whatsapp.send_whatsapp_text = original_send
            whatsapp.send_whatsapp_list = original_send_list
            whatsapp.user_states.pop("08123456789", None)
            db.close()

    def test_whatsapp_panduan_commands_resend_guide_without_clearing_form_state(self):
        db = TestingSessionLocal()
        employee = Employee(nama="Panduan WhatsApp", profile_code="OFFICE", active=True)
        db.add(employee)
        db.flush()
        db.add(whatsapp.models.WhatsAppIdentity(
            phone_number="628123456789",
            employee_id=employee.id,
            status="ACTIVE",
        ))
        db.commit()
        sent = []

        async def fake_send(to_number, message):
            sent.append(message)
            return {"mocked": True}

        original_send = whatsapp.send_whatsapp_text
        whatsapp.send_whatsapp_text = fake_send
        whatsapp.user_states["08123456789"] = {
            "step": "pilih_kategori",
            "data": {"preserve": True},
        }
        try:
            asyncio.run(whatsapp.process_whatsapp_incoming(
                "08123456789", "panduan", None, db
            ))
            asyncio.run(whatsapp.process_whatsapp_incoming(
                "08123456789", "/panduan", None, db
            ))

            self.assertEqual(len(sent), 2)
            self.assertEqual(sent[0], sent[1])
            self.assertIn("Panduan HR WhatsApp", sent[0])
            self.assertEqual(whatsapp.user_states["08123456789"]["step"], "pilih_kategori")
            self.assertEqual(
                whatsapp.user_states["08123456789"]["data"],
                {"preserve": True},
            )
        finally:
            whatsapp.send_whatsapp_text = original_send
            whatsapp.user_states.pop("08123456789", None)
            db.close()

    def test_whatsapp_profile_shows_only_linked_employee_fields_without_nik(self):
        db = TestingSessionLocal()
        employee = Employee(
            nama="Profil WhatsApp",
            nik="3201010101010099",
            profile_code="OFFICE",
            cabang="Jakarta",
            employment_status="PKWT",
            active=True,
        )
        db.add(employee)
        db.flush()
        db.add(whatsapp.models.WhatsAppIdentity(
            phone_number="628123456789",
            employee_id=employee.id,
            status="ACTIVE",
        ))
        db.commit()
        sent = []

        async def fake_send(to_number, message):
            sent.append(message)
            return {"mocked": True}

        original_send = whatsapp.send_whatsapp_text
        whatsapp.send_whatsapp_text = fake_send
        whatsapp.user_states.pop("08123456789", None)
        try:
            asyncio.run(whatsapp.process_whatsapp_incoming(
                "08123456789", "", "menu_profile", db
            ))
            self.assertIn("Profil WhatsApp", sent[-1])
            self.assertIn("Jakarta", sent[-1])
            self.assertIn("PKWT", sent[-1])
            self.assertNotIn(employee.nik, sent[-1])
            self.assertNotIn("NIK", sent[-1])

            asyncio.run(whatsapp.process_whatsapp_incoming(
                "08123456789", "profil", None, db
            ))
            self.assertIn("Profil WhatsApp", sent[-1])
            self.assertNotIn(employee.nik, sent[-1])
        finally:
            whatsapp.send_whatsapp_text = original_send
            whatsapp.user_states.pop("08123456789", None)
            db.close()

    def test_whatsapp_cuti_and_izin_menu_routes_to_separate_categories(self):
        db = TestingSessionLocal()
        employee = Employee(
            nama="Menu WhatsApp",
            profile_code="OFFICE",
            active=True,
        )
        db.add(employee)
        db.flush()
        db.add(whatsapp.models.WhatsAppIdentity(
            phone_number="628123456789",
            employee_id=employee.id,
            status="ACTIVE",
        ))
        db.commit()
        sent_lists = []
        original_send_list = whatsapp.send_whatsapp_list
        original_check_quota = whatsapp.check_quota_and_prompt_date

        async def fake_send_list(**kwargs):
            sent_lists.append(kwargs)
            return {"mocked": True}

        async def fake_check_quota(to_number, nama, _db, is_wfl=False):
            whatsapp.user_states[to_number]["step"] = "input_tgl"

        whatsapp.send_whatsapp_list = fake_send_list
        whatsapp.check_quota_and_prompt_date = fake_check_quota
        whatsapp.user_states.pop("08123456789", None)
        try:
            asyncio.run(whatsapp.process_whatsapp_incoming(
                "08123456789", "", "menu_cuti", db
            ))
            self.assertEqual(
                whatsapp.user_states["08123456789"]["data"]["kategori"],
                "CUTI_TAHUNAN",
            )
            self.assertEqual(whatsapp.user_states["08123456789"]["step"], "input_tgl")

            asyncio.run(whatsapp.process_whatsapp_incoming(
                "08123456789", "", "menu_izin", db
            ))
            self.assertEqual(
                whatsapp.user_states["08123456789"]["data"]["allowed_categories"],
                ["SAKIT", "IZIN_TELAT", "IZIN_PULANG_CEPAT", "LAINNYA"],
            )
            self.assertEqual(
                [row["id"] for row in sent_lists[-1]["sections"][0]["rows"]],
                ["kat_SAKIT", "kat_IZIN_TELAT", "kat_IZIN_PULANG_CEPAT", "kat_LAINNYA"],
            )
            asyncio.run(whatsapp.process_whatsapp_incoming(
                "08123456789", "", "kat_IZIN_TELAT", db
            ))
            self.assertEqual(
                whatsapp.user_states["08123456789"]["data"]["kategori"],
                "IZIN_TELAT",
            )
        finally:
            whatsapp.send_whatsapp_list = original_send_list
            whatsapp.check_quota_and_prompt_date = original_check_quota
            whatsapp.user_states.pop("08123456789", None)
            db.close()

    def test_9_whatsapp_locks_all_messages_before_login(self):
        db = TestingSessionLocal()
        sent = []

        async def fake_send(to_number, message):
            sent.append(message)
            return {"mocked": True}

        original_send = whatsapp.send_whatsapp_text
        whatsapp.send_whatsapp_text = fake_send
        try:
            asyncio.run(whatsapp.process_whatsapp_incoming("081111111111", "menu", None, db))
            self.assertIn("belum login", sent[-1])
            self.assertNotIn("step", whatsapp.user_states.get("081111111111", {}))

            sent.clear()
            asyncio.run(whatsapp.process_whatsapp_incoming("081111111111", "cuti", None, db))
            self.assertIn("belum login", sent[-1])

            sent.clear()
            asyncio.run(whatsapp.process_whatsapp_incoming("081111111111", "profil", None, db))
            self.assertIn("belum login", sent[-1])

            sent.clear()
            asyncio.run(whatsapp.process_whatsapp_incoming("081111111111", "login", None, db))
            self.assertEqual(whatsapp.user_states["081111111111"]["step"], "login_waiting_nik")
        finally:
            whatsapp.send_whatsapp_text = original_send
            whatsapp.user_states.pop("081111111111", None)
            db.close()

    def test_10_whatsapp_message_claim_is_idempotent(self):
        db = TestingSessionLocal()
        self.assertTrue(whatsapp.claim_whatsapp_message(db, "wamid.TEST-001"))
        self.assertFalse(whatsapp.claim_whatsapp_message(db, "wamid.TEST-001"))
        self.assertTrue(whatsapp.claim_whatsapp_message(db, "wamid.TEST-002"))
        self.assertEqual(db.query(whatsapp.models.WhatsAppProcessedMessage).count(), 2)
        db.commit()
        db.close()

    def test_11_whatsapp_conversation_state_expires(self):
        db = TestingSessionLocal()
        sent = []

        async def fake_send(to_number, message):
            sent.append(message)
            return {"mocked": True}

        original_send = whatsapp.send_whatsapp_text
        whatsapp.send_whatsapp_text = fake_send
        whatsapp.user_states["628199999999"] = {
            "step": "login_waiting_nik",
            "data": {},
            "expires_at": datetime.datetime.utcnow() - datetime.timedelta(minutes=1),
        }
        try:
            asyncio.run(whatsapp.process_whatsapp_incoming("628199999999", "123", None, db))
            self.assertIn("belum login", sent[-1])
            self.assertNotIn("628199999999", whatsapp.user_states)
        finally:
            whatsapp.send_whatsapp_text = original_send
            whatsapp.user_states.pop("628199999999", None)
            db.close()

    def test_12_hr_can_audit_revoke_and_move_whatsapp_linking(self):
        db = TestingSessionLocal()
        first = Employee(
            nama="Karyawan Link Satu",
            employee_code="JIP-10001",
            nik="3201010101010002",
            profile_code="OFFICE",
            active=True,
        )
        second = Employee(
            nama="Karyawan Link Dua",
            employee_code="JIP-10002",
            nik="3201010101010003",
            profile_code="OFFICE",
            active=False,
        )
        db.add_all([first, second])
        db.commit()
        db.refresh(first)
        db.refresh(second)
        first_id = first.id
        second_id = second.id
        identity = whatsapp.models.WhatsAppIdentity(
            phone_number="628122233344",
            employee_id=first_id,
            status="ACTIVE",
        )
        target_identity = whatsapp.models.WhatsAppIdentity(
            phone_number="628122233355",
            employee_id=second_id,
            status="ACTIVE",
        )
        db.add_all([identity, target_identity])
        db.commit()
        db.refresh(identity)
        identity_id = identity.id
        target_identity_id = target_identity.id
        db.close()

        audit_res = client.get("/employees/whatsapp-audit")
        self.assertEqual(audit_res.status_code, 200)
        self.assertIsInstance(audit_res.json(), list)

        move_res = client.patch(
            f"/employees/whatsapp-identities/{identity_id}/move",
            json={"employee_code": "JIP-10002", "reason": "Nomor dikonfirmasi HR sebagai milik karyawan kedua"},
        )
        self.assertEqual(move_res.status_code, 409)

        db = TestingSessionLocal()
        db.query(whatsapp.models.WhatsAppIdentity).filter(
            whatsapp.models.WhatsAppIdentity.id == target_identity_id
        ).delete()
        db.commit()
        db.close()

        move_res = client.patch(
            f"/employees/whatsapp-identities/{identity_id}/move",
            json={"employee_code": "JIP-10002", "reason": "Nomor dikonfirmasi HR sebagai milik karyawan kedua"},
        )
        self.assertEqual(move_res.status_code, 200)
        self.assertEqual(move_res.json()["employee_id"], second_id)
        self.assertEqual(move_res.json()["employee_code"], "JIP-10002")
        self.assertEqual(client.patch(
            f"/employees/whatsapp-identities/{identity_id}/move",
            json={"employee_code": "JIP-UNKNOWN", "reason": "Kode tidak terdaftar"},
        ).status_code, 404)
        self.assertEqual(client.patch(
            f"/employees/whatsapp-identities/{identity_id}/move",
            json={"reason": "Kode karyawan tidak diberikan"},
        ).status_code, 400)

        revoke_res = client.patch(f"/employees/whatsapp-identities/{identity_id}/revoke")
        self.assertEqual(revoke_res.status_code, 200)
        self.assertEqual(revoke_res.json()["new_status"], "REVOKED")

        db = TestingSessionLocal()
        events = {
            row.event_type
            for row in db.query(whatsapp.models.WhatsAppAuditLog).all()
        }
        self.assertIn("HR_MOVE", events)
        self.assertIn("HR_REVOKE", events)
        db.close()


class TestEmployeeWorkbookImports(unittest.TestCase):

    def setUp(self):
        Base.metadata.drop_all(bind=engine)
        Base.metadata.create_all(bind=engine)
        db = TestingSessionLocal()
        db.add(Profile(
            code="OFFICE",
            nama="Office",
            hari_kerja=["Senin", "Selasa", "Rabu", "Kamis", "Jumat"],
        ))
        db.commit()
        db.close()

    @staticmethod
    def _upload(workbook, filename):
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return {"file": (filename, output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    @staticmethod
    def _master_workbook():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Master ID"
        sheet.append([
            "ID Karyawan", "Kode Cabang", "No Urut", "Nama", "NIK",
            "Tanggal Join", "UM (Uang Makan)", "Jabatan", "Lokasi Kerja",
            "Sheet Asal", "Status (Bulanan/Harian)",
        ])
        sheet.append([
            "JIP-010001", "01", "0001", "Current Employee", "3201010101010001",
            datetime.date(2024, 1, 15), 65000, "Admin", "Gudang A", "Gudang A", "Bulanan",
        ])
        sheet.append([
            "JIP-010002", "01", "0002", "New Employee", "3201010101010002",
            datetime.date(2025, 2, 15), 70000, "Staff", "Gudang A", "Gudang A", "Harian",
        ])
        return workbook

    def test_master_dry_run_commit_and_review_new_employee(self):
        db = TestingSessionLocal()
        employee = Employee(
            nama="Old Employee",
            nik="3201010101010001",
            profile_code="OFFICE",
            employment_status="PKWT",
            active=False,
            cabang="OLD BRANCH",
        )
        db.add(employee)
        db.add(LeaveRequest(
            nama="Old Employee",
            kategori="SAKIT",
            tanggal_mulai=datetime.date(2025, 1, 1),
            tanggal_selesai=datetime.date(2025, 1, 1),
            status="APPROVED",
        ))
        db.commit()
        employee_id = employee.id
        db.close()

        workbook = self._master_workbook()
        dry = client.post(
            "/employees/import/master?dry_run=true",
            files=self._upload(workbook, "master.xlsx"),
        )
        self.assertEqual(dry.status_code, 200)
        self.assertEqual(dry.json()["total_updated"], 1)
        self.assertEqual(dry.json()["total_new_pending_review"], 1)
        self.assertEqual(dry.json()["total_conflicts"], 0)

        db = TestingSessionLocal()
        unchanged = db.query(Employee).filter(Employee.id == employee_id).one()
        self.assertIsNone(unchanged.employee_code)
        self.assertEqual(unchanged.cabang, "OLD BRANCH")
        self.assertEqual(db.query(EmployeeImportCandidate).count(), 0)
        db.close()

        applied = client.post(
            "/employees/import/master?dry_run=false",
            files=self._upload(workbook, "master.xlsx"),
        )
        self.assertEqual(applied.status_code, 200)
        self.assertEqual(applied.json()["total_updated"], 1)
        self.assertEqual(applied.json()["total_new_pending_review"], 1)

        db = TestingSessionLocal()
        updated = db.query(Employee).filter(Employee.id == employee_id).one()
        self.assertEqual(updated.employee_code, "JIP-010001")
        self.assertEqual(updated.nama, "Current Employee")
        self.assertEqual(updated.cabang, "Gudang A")
        self.assertTrue(updated.active)
        self.assertEqual(updated.profile_code, "OFFICE")
        self.assertEqual(updated.employment_status, "PKWT")
        self.assertEqual(db.query(LeaveRequest).one().nama, "Current Employee")
        candidate = db.query(EmployeeImportCandidate).filter(
            EmployeeImportCandidate.employee_code == "JIP-010002"
        ).one()
        candidate_id = candidate.id
        self.assertEqual(candidate.status, "PENDING")
        db.close()

        approved = client.post(
            f"/employees/import/master/pending/{candidate_id}/approve",
            json={"profile_code": "OFFICE", "employment_status": "PKWTT"},
        )
        self.assertEqual(approved.status_code, 200)
        self.assertTrue(approved.json()["active"])
        self.assertEqual(approved.json()["employee_code"], "JIP-010002")
        self.assertEqual(approved.json()["payroll_status"], "Harian")

        updated = client.put(
            f"/employees/{employee_id}",
            json={
                "nama": "Current Employee",
                "profile_code": "OFFICE",
                "employment_status": "PKWT",
                "join_date": "2024-01-15",
                "cabang": "Gudang A",
                "active": True,
            },
        )
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.json()["employee_code"], "JIP-010001")
        db = TestingSessionLocal()
        preserved = db.query(Employee).filter(Employee.id == employee_id).one()
        self.assertEqual(preserved.jabatan, "Admin")
        self.assertEqual(preserved.payroll_status, "Bulanan")
        db.close()

    def test_current_employee_template_uses_hr_fields_without_pay_or_contract_end_columns(self):
        response = client.get("/employees/template")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content), data_only=True)
        self.assertEqual(workbook.sheetnames, ["Data Karyawan", "Referensi", "Petunjuk"])
        headers = [cell.value for cell in workbook["Data Karyawan"][1]]
        self.assertIn("Kode Karyawan", headers)
        self.assertIn("Kode Profil", headers)
        self.assertIn("Status Kepegawaian", headers)
        self.assertNotIn("Gaji Pokok", headers)
        self.assertNotIn("Waktu Berakhir", headers)
        profile_codes = [workbook["Referensi"].cell(row, 1).value for row in range(2, workbook["Referensi"].max_row + 1)]
        self.assertIn("OFFICE", profile_codes)

    def test_new_import_format_requires_code_and_keeps_pkwt_inactive_until_contract(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data Karyawan"
        sheet.append([
            "Kode Karyawan", "Nama Karyawan", "NIK", "Cabang", "Kode Profil",
            "Status Kepegawaian", "Tanggal Masuk", "Status Aktif", "Jabatan", "ID Mesin Absensi",
        ])
        sheet.append([
            "EMP-TEST-001", "Pegawai Uji", "9999999999999999", "Cabang Uji",
            "OFFICE", "PKWT", datetime.date(2025, 1, 2), "Aktif", "Staff", "M-001",
        ])
        response = client.post(
            "/employees/import",
            files=self._upload(workbook, "master_karyawan.xlsx"),
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["inserted"], 1)
        self.assertEqual(response.json()["contract_required"], ["EMP-TEST-001"])
        db = TestingSessionLocal()
        employee = db.query(Employee).filter(Employee.employee_code == "EMP-TEST-001").one()
        self.assertEqual(employee.nik, "9999999999999999")
        self.assertEqual(employee.employment_status, "PKWT")
        self.assertFalse(employee.active)
        self.assertEqual(employee.contracts, [])
        db.close()

        legacy = Workbook()
        old_sheet = legacy.active
        old_sheet.title = "HO"
        old_sheet.append(["Status", "Nama", "Gaji Pokok", "Waktu Berakhir"])
        old_sheet.append(["Office", "Pegawai Lama", "", ""])
        rejected = client.post(
            "/employees/import",
            files=self._upload(legacy, "template_lama.xlsx"),
        )
        self.assertEqual(rejected.status_code, 400)

    def test_master_update_matches_by_employee_code_and_does_not_touch_contract_history(self):
        db = TestingSessionLocal()
        employee = Employee(
            employee_code="EMP-UPDATE-001",
            nama="Pegawai Uji",
            nik="9999999999999998",
            profile_code="OFFICE",
            employment_status="PKWTT",
            join_date=datetime.date(2022, 1, 1),
            active=True,
            cabang="Cabang Lama",
        )
        db.add(employee)
        db.commit()
        db.close()

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Pembaruan Master"
        sheet.append([
            "Kode Karyawan", "Nama Karyawan", "NIK", "Cabang", "Kode Profil",
            "Status Kepegawaian", "Tanggal Masuk", "Status Aktif", "Jabatan", "ID Mesin Absensi",
        ])
        sheet.append([
            "EMP-UPDATE-001", "Pegawai Uji", "9999999999999997", "Cabang Baru",
            "OFFICE", "PKWT", datetime.date(2023, 2, 3), "Aktif", "Admin", "M-002",
        ])
        response = client.post(
            "/employees/import/nik",
            files=self._upload(workbook, "update_master.xlsx"),
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["total_updated"], 1)
        self.assertEqual(response.json()["detail"]["contract_required"], ["EMP-UPDATE-001"])
        db = TestingSessionLocal()
        updated = db.query(Employee).filter(Employee.employee_code == "EMP-UPDATE-001").one()
        self.assertEqual(updated.nik, "9999999999999997")
        self.assertEqual(updated.cabang, "Cabang Baru")
        self.assertEqual(updated.employment_status, "PKWT")
        self.assertFalse(updated.active)
        self.assertEqual(db.query(EmploymentContract).filter(EmploymentContract.employee_id == updated.id).count(), 0)
        db.close()

    def test_phl_leave_requests_are_blocked_without_changing_history(self):
        db = TestingSessionLocal()
        employee = Employee(
            employee_code="EMP-PHL-001",
            nama="Karyawan PHL Uji",
            profile_code="OFFICE",
            employment_status="PHL",
            active=True,
        )
        db.add(employee)
        db.add(LeaveRequest(
            nama="Karyawan PHL Uji",
            kategori="SAKIT",
            tanggal_mulai=datetime.date(2024, 1, 2),
            tanggal_selesai=datetime.date(2024, 1, 2),
            status="APPROVED",
        ))
        db.commit()
        db.close()

        response = client.post("/leaves", json={
            "nama": "Karyawan PHL Uji",
            "kategori": "SAKIT",
            "tanggal_mulai": "2026-10-10",
            "tanggal_selesai": "2026-10-10",
            "alasan": "Uji",
        })
        self.assertEqual(response.status_code, 400)
        self.assertIn("PHL", response.json()["detail"])

        remote = client.post("/leaves", json={
            "nama": "Karyawan PHL Uji",
            "kategori": "WORK_FROM_LOCATION",
            "tanggal_mulai": "2026-10-12",
            "tanggal_selesai": "2026-10-12",
            "alasan": "Tugas luar",
            "tipe_absensi": "remote_work",
            "location_cabang": "Cabang Uji",
        })
        self.assertEqual(remote.status_code, 200)
        db = TestingSessionLocal()
        self.assertEqual(db.query(LeaveRequest).count(), 2)
        historic = db.query(LeaveRequest).filter(LeaveRequest.tanggal_mulai == datetime.date(2024, 1, 2)).one()
        self.assertEqual(historic.status, "APPROVED")
        db.close()

    def test_phl_whatsapp_cuti_and_izin_are_blocked(self):
        db = TestingSessionLocal()
        employee = Employee(
            employee_code="EMP-PHL-WA",
            nama="PHL WhatsApp",
            profile_code="OFFICE",
            employment_status="PHL",
            active=True,
        )
        db.add(employee)
        db.flush()
        db.add(whatsapp.models.WhatsAppIdentity(
            phone_number="628123456789",
            employee_id=employee.id,
            status="ACTIVE",
        ))
        db.commit()
        sent = []
        original_send = whatsapp.send_whatsapp_text

        async def fake_send(to_number, message):
            sent.append(message)
            return {"mocked": True}

        whatsapp.send_whatsapp_text = fake_send
        whatsapp.user_states.pop("08123456789", None)
        try:
            asyncio.run(whatsapp.process_whatsapp_incoming("08123456789", "", "menu_cuti", db))
            self.assertIn("PHL", sent[-1])
            self.assertNotIn("08123456789", whatsapp.user_states)
            asyncio.run(whatsapp.process_whatsapp_incoming("08123456789", "", "menu_izin", db))
            self.assertIn("PHL", sent[-1])
            self.assertNotIn("08123456789", whatsapp.user_states)
        finally:
            whatsapp.send_whatsapp_text = original_send
            whatsapp.user_states.pop("08123456789", None)
            db.close()

    def test_phl_attendance_review_only_allows_unpaid_leave(self):
        db = TestingSessionLocal()
        employee = Employee(
            employee_code="EMP-PHL-REVIEW",
            nama="PHL Review",
            profile_code="OFFICE",
            employment_status="PHL",
            active=True,
        )
        db.add(employee)
        db.flush()
        review = whatsapp.models.AttendanceReview(
            employee_id=employee.id,
            attendance_date=datetime.date(2026, 10, 11),
            status="PENDING",
        )
        db.add(review)
        db.commit()
        review_id = review.id
        db.close()

        rejected = client.patch(
            f"/attendance/reviews/{review_id}",
            json={"decision": "ANNUAL_LEAVE"},
        )
        self.assertEqual(rejected.status_code, 400)
        db = TestingSessionLocal()
        self.assertEqual(db.query(whatsapp.models.AttendanceReview).filter_by(id=review_id).one().status, "PENDING")
        db.close()

        accepted = client.patch(
            f"/attendance/reviews/{review_id}",
            json={"decision": "UNPAID_LEAVE"},
        )
        self.assertEqual(accepted.status_code, 200)
        self.assertEqual(accepted.json()["decision"], "UNPAID_LEAVE")
        db = TestingSessionLocal()
        leave = db.query(LeaveRequest).one()
        self.assertEqual(leave.kategori, "UNPAID_LEAVE")
        self.assertEqual(leave.status, "APPROVED")
        db.close()

    def test_recap_import_quarantines_missing_ids_and_is_idempotent(self):
        db = TestingSessionLocal()
        db.add(Employee(
            employee_code="JIP-020001",
            nama="Recap Employee",
            nik="3201010101010003",
            profile_code="OFFICE",
            active=True,
        ))
        db.commit()
        db.close()

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Rekap Kehadiran"
        sheet.append(["No", "ID", "NIK", "Cabang", "Nama", "Status", "Tanggal", "Keterangan (S/I/CS)"])
        sheet.append([1, "JIP-020001", "3201010101010003", "Kby. Lama", "Recap Employee", "Bulanan", datetime.date(2026, 1, 1), "C"])
        sheet.append([2, "JIP-020001", "3201010101010003", "Kby. Lama", "Recap Employee", "Bulanan", datetime.date(2026, 1, 2), "CS"])
        sheet.append([3, "JIP-020001", "3201010101010003", "Kby. Lama", "Recap Employee", "Bulanan", datetime.date(2026, 1, 3), "S"])
        sheet.append([4, "JIP-020001", "3201010101010003", "Kby. Lama", "Recap Employee", "Bulanan", datetime.date(2026, 1, 4), "T"])
        sheet.append([5, None, "3201010101010003", "Kby. Lama", "Recap Employee", "Bulanan", datetime.date(2026, 1, 5), "C"])

        dry = client.post(
            "/employees/import/recap?dry_run=true",
            files=self._upload(workbook, "recap.xlsx"),
        )
        self.assertEqual(dry.status_code, 200)
        self.assertEqual(dry.json()["total_to_import"], 3)
        self.assertEqual(dry.json()["total_t_ignored"], 1)
        self.assertEqual(dry.json()["total_conflicts"], 1)
        db = TestingSessionLocal()
        self.assertEqual(db.query(LeaveRequest).count(), 0)
        db.close()

        applied = client.post(
            "/employees/import/recap?dry_run=false",
            files=self._upload(workbook, "recap.xlsx"),
        )
        self.assertEqual(applied.status_code, 200)
        self.assertEqual(applied.json()["total_to_import"], 3)

        repeated = client.post(
            "/employees/import/recap?dry_run=false",
            files=self._upload(workbook, "recap.xlsx"),
        )
        self.assertEqual(repeated.status_code, 200)
        self.assertEqual(repeated.json()["total_to_import"], 0)
        self.assertEqual(repeated.json()["total_duplicate_or_existing"], 3)

        db = TestingSessionLocal()
        self.assertEqual(db.query(LeaveRequest).count(), 3)
        self.assertEqual(db.query(LeaveRequest).filter(LeaveRequest.kategori == "SAKIT").count(), 1)
        self.assertEqual(db.query(LeaveBalanceLedger).filter(LeaveBalanceLedger.entry_type == "USED").count(), 2)
        db.close()

    def test_missing_employee_deactivation_requires_explicit_apply_and_never_deletes(self):
        db = TestingSessionLocal()
        db.add_all([
            Employee(
                employee_code="JIP-010001",
                nama="Current Employee",
                nik="3201010101010001",
                profile_code="OFFICE",
                active=True,
            ),
            Employee(
                employee_code="JIP-OLD0001",
                nama="Former Employee",
                nik="3201010101010009",
                profile_code="OFFICE",
                active=True,
            ),
        ])
        db.commit()
        former_id = db.query(Employee).filter(Employee.employee_code == "JIP-OLD0001").one().id
        db.close()

        preview = client.post(
            "/employees/import/master?dry_run=true&deactivate_missing=true",
            files=self._upload(self._master_workbook(), "master.xlsx"),
        )
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.json()["total_missing_from_master"], 1)
        db = TestingSessionLocal()
        self.assertTrue(db.query(Employee).filter(Employee.id == former_id).one().active)
        db.close()

        applied = client.post(
            "/employees/import/master?dry_run=false&deactivate_missing=true",
            files=self._upload(self._master_workbook(), "master.xlsx"),
        )
        self.assertEqual(applied.status_code, 200)
        db = TestingSessionLocal()
        former = db.query(Employee).filter(Employee.id == former_id).one()
        self.assertFalse(former.active)
        self.assertEqual(former.employee_code, "JIP-OLD0001")
        db.close()


if __name__ == "__main__":
    unittest.main()
