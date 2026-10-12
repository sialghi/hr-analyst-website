from collections.abc import Generator

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import auth, models
from app.database import Base, get_db
from app.main import app
from app.shared_store import _file_store


@pytest.fixture
def process_client(tmp_path) -> Generator[tuple[TestClient, sessionmaker, str], None, None]:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    test_session_factory = sessionmaker(autocommit=False, autoflush=False, bind=engine)

    def override_get_db():
        db = test_session_factory()
        try:
            yield db
        finally:
            db.close()

    def override_auth():
        return models.User(
            id=1,
            nama="Synthetic HR",
            email="synthetic.hr@example.com",
            role=models.RoleEnum.hr_master,
        )

    original_overrides = app.dependency_overrides.copy()
    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[auth.get_current_user] = override_auth

    file_id = "synthetic-adjustment-report"
    report_path = tmp_path / "report.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "2. Summary Overview"
    worksheet.cell(4, 2, "Nama")
    worksheet.cell(4, 23, "Total Bonus Lain-lain (Rp)")
    worksheet.cell(4, 24, "Total Potongan Lain-lain (Rp)")
    worksheet.cell(4, 31, "Total Bersih (Rp)")
    worksheet.cell(5, 2, "Karyawan QA")
    worksheet.cell(5, 23, 100)
    worksheet.cell(5, 24, 50)
    worksheet.cell(5, 31, "=AD5-U5-V5-X5+W5+Y5")
    workbook.save(report_path)

    report_row = {
        "Nama": "Karyawan QA",
        "Total Bonus Lain-lain (Rp)": 100,
        "Total Potongan Lain-lain (Rp)": 50,
        "Total Uang Makan Akhir (Rp)": 1000,
        "Potongan BPJS Kesehatan (Rp)": 0,
        "Potongan BPJS TK (Rp)": 0,
        "Total Bonus Lembur (Rp)": 0,
        "Total Bersih (Rp)": 1050,
    }
    result_data = {
        "kpis": {"total_bersih": 1050},
        "sheets": [{
            "key": "summary",
            "data": {
                "headers": [
                    "Nama",
                    "Total Bonus Lain-lain (Rp)",
                    "Total Potongan Lain-lain (Rp)",
                    "Total Bersih (Rp)",
                ],
                "rows": [report_row],
            },
        }],
    }
    _file_store[file_id] = {
        "path": str(report_path),
        "filename": "report.xlsx",
        "result_data": result_data,
    }

    with test_session_factory() as db:
        db.add(models.EmployeeAdjustment(nama="Karyawan QA", bonus_lain=100, potongan_lain=50))
        db.commit()

    client = TestClient(app)
    try:
        yield client, test_session_factory, file_id
    finally:
        client.close()
        _file_store.pop(file_id, None)
        app.dependency_overrides.clear()
        app.dependency_overrides.update(original_overrides)
        engine.dispose()


def test_adjustment_updates_summary_kpi_and_download_workbook(process_client):
    client, session_factory, file_id = process_client

    bonus_update = client.patch(
        f"/proses/{file_id}/adjustments",
        json={"nama": "Karyawan QA", "bonus_lain": 250, "potongan_lain": 50},
    )
    assert bonus_update.status_code == 200
    assert bonus_update.json()["total_bersih"] == 1200

    deduction_update = client.patch(
        f"/proses/{file_id}/adjustments",
        json={"nama": "Karyawan QA", "bonus_lain": 250, "potongan_lain": 125},
    )
    assert deduction_update.status_code == 200
    assert deduction_update.json()["total_bersih"] == 1125

    result_data = _file_store[file_id]["result_data"]
    row = result_data["sheets"][0]["data"]["rows"][0]
    assert row["Total Bonus Lain-lain (Rp)"] == 250
    assert row["Total Potongan Lain-lain (Rp)"] == 125
    assert row["Total Bersih (Rp)"] == 1125
    assert result_data["kpis"]["total_bersih"] == 1125

    workbook = load_workbook(_file_store[file_id]["path"], data_only=False)
    worksheet = workbook["2. Summary Overview"]
    assert worksheet.cell(5, 23).value == 250
    assert worksheet.cell(5, 24).value == 125
    assert worksheet.cell(5, 31).value == "=AD5-U5-V5-X5+W5+Y5"
    assert workbook.calculation.fullCalcOnLoad is True
    workbook.close()

    with session_factory() as db:
        saved = db.query(models.EmployeeAdjustment).filter_by(nama="Karyawan QA").one()
        assert saved.bonus_lain == 250
        assert saved.potongan_lain == 125


def test_adjustment_update_rejects_unknown_report_or_employee_without_saving(process_client):
    client, session_factory, file_id = process_client

    missing_report = client.patch(
        "/proses/not-a-real-file/adjustments",
        json={"nama": "Karyawan QA", "bonus_lain": 999, "potongan_lain": 0},
    )
    missing_employee = client.patch(
        f"/proses/{file_id}/adjustments",
        json={"nama": "Tidak Ada", "bonus_lain": 999, "potongan_lain": 0},
    )

    assert missing_report.status_code == 404
    assert missing_employee.status_code == 404
    with session_factory() as db:
        saved = db.query(models.EmployeeAdjustment).filter_by(nama="Karyawan QA").one()
        assert saved.bonus_lain == 100
        assert saved.potongan_lain == 50
