# -*- coding: utf-8 -*-
"""Parsing and normalization helpers for employee workbooks."""
import datetime
import io
import re
import csv
from collections import Counter

import openpyxl

from .branches import branch_code_for, employee_id_serial, validate_new_employee_id


def normalize_text(value) -> str:
    return " ".join(str(value or "").split()).casefold()


def normalize_header(value) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(value or "").casefold()).split())


def _is_blank(value) -> bool:
    return value is None or (
        isinstance(value, str)
        and value.strip().casefold() in ("", "-", "none", "nan", "n/a")
    )


def _raw_source_value(value):
    if value is None:
        return None
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.isoformat()
    return str(value)


def clean_nik(value):
    if _is_blank(value):
        return None
    text = str(value).strip().replace(" ", "")
    if text.endswith(".0"):
        text = text[:-2]
    if not text:
        return None
    return text if len(text) == 16 and text.isdigit() else False


def parse_date(value):
    if _is_blank(value):
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return False


def _parse_active(value):
    if _is_blank(value):
        return True
    normalized = str(value).strip().casefold()
    if normalized in ("aktif", "active", "ya", "yes", "true", "1"):
        return True
    if normalized in ("nonaktif", "non-aktif", "inactive", "tidak", "no", "false", "0"):
        return False
    return False


def parse_employee_workbook(content, filename):
    """Parse the current employee import format; legacy layouts are intentionally rejected."""
    aliases = {
        "new_employee_id": ("new employee id", "id karyawan", "id karyawan baru"),
        "nama": ("nama karyawan", "nama", "employee name"),
        "nik": ("nik", "nik ktp"),
        "cabang": ("cabang", "lokasi kerja"),
        "profile_code": ("kode profil", "profil", "profil divisi"),
        "employment_status": ("status kepegawaian",),
        "join_date": ("tanggal masuk", "tanggal join"),
        "active": ("status aktif",),
        "jabatan": ("jabatan",),
        "id_mesin": ("id mesin absensi", "id mesin absen"),
    }
    extension = str(filename or "").lower().rsplit(".", 1)[-1]
    if extension == "csv":
        try:
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
            source_headers = list(reader.fieldnames or [])
            source_rows = list(reader)
            header_row = 1
        except (UnicodeDecodeError, csv.Error) as exc:
            raise ValueError("File CSV tidak valid; gunakan UTF-8.") from exc
        sheet_name = None
        normalized_headers = {
            normalize_header(header): index
            for index, header in enumerate(source_headers)
            if header is not None
        }
    elif extension in ("xlsx", "xlsm"):
        workbook = _read_workbook(content)
        sheet = None
        header_row = None
        normalized_headers = None
        for candidate in workbook.worksheets:
            if candidate.title.casefold() not in ("data karyawan", "data karyawan "):
                continue
            for row_number, values in enumerate(
                candidate.iter_rows(min_row=1, max_row=min(candidate.max_row, 10), values_only=True),
                start=1,
            ):
                headers = {
                    normalize_header(value): index
                    for index, value in enumerate(values)
                    if value is not None
                }
                if all(name in headers for name in ("id karyawan", "nama karyawan", "kode profil", "status kepegawaian")):
                    sheet, header_row, normalized_headers = candidate, row_number, headers
                    break
            if sheet:
                break
        if sheet is None:
            raise ValueError("Format workbook tidak dikenali. Gunakan sheet 'Data Karyawan' dari template terbaru.")
        source_headers = next(
            sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True)
        )
        source_rows = sheet.iter_rows(min_row=header_row + 1, values_only=True)
        sheet_name = sheet.title
    else:
        raise ValueError("Format file tidak didukung. Gunakan .xlsx, .xlsm, atau .csv dengan format template terbaru.")

    column_indexes = {}
    for field, candidates in aliases.items():
        for candidate in candidates:
            if candidate in normalized_headers:
                column_indexes[field] = normalized_headers[candidate]
                break
    for required in ("new_employee_id", "nama", "cabang", "profile_code", "employment_status"):
        if required not in column_indexes:
            label = {
                "new_employee_id": "ID Karyawan",
                "nama": "Nama Karyawan",
                "cabang": "Cabang",
                "profile_code": "Kode Profil",
                "employment_status": "Status Kepegawaian",
            }[required]
            raise ValueError(f"Kolom wajib '{label}' tidak ditemukan pada file.")

    parsed_rows = []
    errors = []
    for row_number, values in enumerate(source_rows, start=header_row + 1):
        if isinstance(values, dict):
            if not any(not _is_blank(value) for value in values.values()):
                continue
            values = [values.get(header) for header in source_headers]
        elif not any(not _is_blank(value) for value in values):
            continue

        def value_for(field):
            index = column_indexes.get(field)
            return values[index] if index is not None and index < len(values) else None

        code_raw = value_for("new_employee_id")
        code = "" if _is_blank(code_raw) else str(code_raw).strip().upper()
        name_raw = value_for("nama")
        name = "" if _is_blank(name_raw) else " ".join(str(name_raw).split())
        branch_raw = value_for("cabang")
        branch = "" if _is_blank(branch_raw) else " ".join(str(branch_raw).split())
        profile_raw = value_for("profile_code")
        profile_code = "" if _is_blank(profile_raw) else str(profile_raw).strip()
        status_raw = value_for("employment_status")
        employment_status = "" if _is_blank(status_raw) else str(status_raw).strip().upper()
        if employment_status in ("TETAP", "KARYAWAN TETAP"):
            employment_status = "PKWTT"
        row_errors = []
        if not code or not name or not branch or not profile_code:
            row_errors.append("ID Karyawan, Nama Karyawan, Cabang, dan Kode Profil wajib diisi.")
        if employment_status not in ("PKWTT", "PKWT", "PHL"):
            row_errors.append("Status Kepegawaian harus PKWTT, PKWT, atau PHL.")

        nik = clean_nik(value_for("nik"))
        if nik is False:
            row_errors.append("NIK harus 16 digit atau dikosongkan.")
        join_date = parse_date(value_for("join_date"))
        if join_date is False:
            row_errors.append("Tanggal Masuk tidak dikenali.")
        active_raw = value_for("active")
        active = _parse_active(active_raw)
        if not _is_blank(active_raw) and str(active_raw).strip().casefold() not in (
            "aktif", "active", "ya", "yes", "true", "1",
            "nonaktif", "non-aktif", "inactive", "tidak", "no", "false", "0",
        ):
            row_errors.append("Status Aktif harus Aktif atau Nonaktif.")
        if row_errors:
            errors.append({"row": row_number, "new_employee_id": code or None, "reason": " ".join(row_errors)})
            continue
        parsed_rows.append({
            "source_row": row_number,
            "new_employee_id": code,
            "nama": name,
            "nik": nik,
            "cabang": branch,
            "profile_code": profile_code,
            "employment_status": employment_status,
            "join_date": join_date,
            "active": active,
            "active_provided": not _is_blank(active_raw),
            "jabatan": None if _is_blank(value_for("jabatan")) else str(value_for("jabatan")).strip(),
            "id_mesin": None if _is_blank(value_for("id_mesin")) else str(value_for("id_mesin")).strip(),
        })

    code_counts = Counter(row["new_employee_id"] for row in parsed_rows)
    serial_counts = Counter(
        serial
        for row in parsed_rows
        if (serial := employee_id_serial(row["new_employee_id"])) is not None
    )
    nik_counts = Counter(row["nik"] for row in parsed_rows if row["nik"])
    for row in parsed_rows:
        reasons = []
        if code_counts[row["new_employee_id"]] > 1:
            reasons.append("ID Karyawan duplikat di file.")
        serial = employee_id_serial(row["new_employee_id"])
        if serial is not None and serial_counts[serial] > 1:
            reasons.append("Nomor urut ID Karyawan harus unik lintas cabang di file.")
        if row["nik"] and nik_counts[row["nik"]] > 1:
            reasons.append("NIK duplikat di file.")
        if reasons:
            errors.append({
                "row": row["source_row"],
                "new_employee_id": row["new_employee_id"],
                "reason": " ".join(reasons),
            })
    error_rows = {error["row"] for error in errors}
    parsed_rows = [row for row in parsed_rows if row["source_row"] not in error_rows]
    if not parsed_rows and not errors:
        raise ValueError("Workbook tidak berisi baris karyawan.")
    return parsed_rows, errors


def parse_money(value):
    if _is_blank(value):
        return None
    if isinstance(value, (int, float)):
        return int(value) if value >= 0 and float(value).is_integer() else False
    text = str(value).strip()
    if not text:
        return None
    if re.fullmatch(r"\d+\.0+", text):
        return int(float(text))
    digits = re.sub(r"[.,\s]", "", text)
    return int(digits) if digits.isdigit() else False


def _read_workbook(content):
    try:
        return openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("File Excel tidak valid atau rusak.") from exc


def _find_sheet_and_header(workbook, required_headers, preferred_sheet=None):
    sheets = list(workbook.worksheets)
    if preferred_sheet:
        sheets.sort(key=lambda sheet: sheet.title.casefold() != preferred_sheet.casefold())
    for sheet in sheets:
        for row_number, cells in enumerate(
            sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 10), values_only=True),
            start=1,
        ):
            headers = {
                normalize_header(value): index
                for index, value in enumerate(cells)
                if value is not None
            }
            if all(required in headers for required in required_headers):
                return sheet, row_number, headers
    raise ValueError("Format workbook tidak dikenali; kolom wajib tidak ditemukan.")


def parse_master_workbook(content):
    workbook = _read_workbook(content)
    sheet, header_row, headers = _find_sheet_and_header(
        workbook,
        ("id karyawan", "nama", "lokasi kerja"),
        preferred_sheet="Master ID",
    )
    aliases = {
        "new_employee_id": ("id karyawan",),
        "nama": ("nama",),
        "nik": ("nik",),
        "join_date": ("tanggal join", "tanggal masuk"),
        "uang_makan_override": ("um uang makan", "uang makan"),
        "jabatan": ("jabatan",),
        "cabang": ("lokasi kerja",),
        "payroll_status": ("status bulanan harian",),
    }
    column_indexes = {}
    for field, candidates in aliases.items():
        for candidate in candidates:
            if candidate in headers:
                column_indexes[field] = headers[candidate]
                break

    parsed_rows = []
    conflicts = []
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        if not any(value is not None for value in values):
            continue
        def value_for(field):
            index = column_indexes.get(field)
            return values[index] if index is not None and index < len(values) else None

        new_employee_id = "" if _is_blank(value_for("new_employee_id")) else str(value_for("new_employee_id")).strip()
        nama = "" if _is_blank(value_for("nama")) else " ".join(str(value_for("nama")).split())
        cabang = "" if _is_blank(value_for("cabang")) else " ".join(str(value_for("cabang")).split())
        if not new_employee_id or not nama or not cabang:
            conflicts.append({"row": row_number, "new_employee_id": new_employee_id or None, "reason": "ID Karyawan, Nama, dan Lokasi Kerja wajib diisi."})
            continue
        if not branch_code_for(cabang):
            conflicts.append({
                "row": row_number,
                "new_employee_id": new_employee_id,
                "reason": f"Cabang '{cabang}' tidak termasuk daftar sembilan cabang yang diizinkan.",
            })
            continue
        if not validate_new_employee_id(new_employee_id, cabang):
            conflicts.append({
                "row": row_number,
                "new_employee_id": new_employee_id,
                "reason": "Awalan ID Karyawan tidak sesuai dengan kode cabang.",
            })
            continue

        nik = clean_nik(value_for("nik"))
        join_date = parse_date(value_for("join_date"))
        allowance = parse_money(value_for("uang_makan_override"))
        payroll_value = value_for("payroll_status")
        payroll_status = None if _is_blank(payroll_value) else " ".join(str(payroll_value).split()).title()
        row_errors = []
        if nik is False:
            row_errors.append("NIK harus 16 digit atau dikosongkan.")
        if join_date is False:
            row_errors.append("Tanggal Join tidak dikenali.")
        if allowance is False:
            row_errors.append("Nilai Uang Makan tidak valid.")
        if payroll_status and payroll_status not in ("Bulanan", "Harian"):
            row_errors.append("Status harus Bulanan atau Harian.")
        if row_errors:
            conflicts.append({"row": row_number, "new_employee_id": new_employee_id, "reason": " ".join(row_errors)})
            continue
        parsed_rows.append({
            "source_row": row_number,
            "new_employee_id": new_employee_id,
            "nama": nama,
            "nik": nik,
            "join_date": join_date,
            "uang_makan_override": allowance,
            "jabatan": None if _is_blank(value_for("jabatan")) else str(value_for("jabatan")).strip(),
            "cabang": cabang,
            "payroll_status": payroll_status,
        })

    code_counts = Counter(row["new_employee_id"] for row in parsed_rows)
    serial_counts = Counter(
        serial
        for row in parsed_rows
        if (serial := employee_id_serial(row["new_employee_id"])) is not None
    )
    nik_counts = Counter(row["nik"] for row in parsed_rows if row["nik"])
    name_counts = Counter(normalize_text(row["nama"]) for row in parsed_rows)
    valid_rows = []
    for row in parsed_rows:
        reasons = []
        if code_counts[row["new_employee_id"]] > 1:
            reasons.append("ID Karyawan duplikat di workbook.")
        serial = employee_id_serial(row["new_employee_id"])
        if serial is not None and serial_counts[serial] > 1:
            reasons.append("Nomor urut ID Karyawan harus unik lintas cabang di workbook.")
        if row["nik"] and nik_counts[row["nik"]] > 1:
            reasons.append("NIK duplikat di workbook.")
        if name_counts[normalize_text(row["nama"])] > 1:
            reasons.append("Nama duplikat; pemrosesan absensi saat ini masih mencocokkan nama.")
        if reasons:
            conflicts.append({
                "row": row["source_row"],
                "new_employee_id": row["new_employee_id"],
                "reason": " ".join(reasons),
            })
        else:
            valid_rows.append(row)
    return valid_rows, conflicts


def parse_recap_workbook(content):
    workbook = _read_workbook(content)
    sheet, header_row, headers = _find_sheet_and_header(
        workbook,
        ("id", "nik", "cabang", "nama", "status", "tanggal"),
        preferred_sheet="Rekap Kehadiran",
    )
    aliases = {
        "new_employee_id": "id",
        "nik": "nik",
        "cabang": "cabang",
        "nama": "nama",
        "status": "status",
        "tanggal": "tanggal",
        "mark": "keterangan s i cs",
    }
    column_indexes = {
        field: headers[header]
        for field, header in aliases.items()
        if header in headers
    }
    if "mark" not in column_indexes:
        raise ValueError("Kolom Keterangan (S/I/CS) tidak ditemukan.")

    rows = []
    conflicts = []
    source_headers = next(
        sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True)
    )
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        if not any(value is not None for value in values):
            continue
        def value_for(field):
            index = column_indexes.get(field)
            return values[index] if index is not None and index < len(values) else None

        new_employee_id = "" if _is_blank(value_for("new_employee_id")) else str(value_for("new_employee_id")).strip()
        nama = "" if _is_blank(value_for("nama")) else " ".join(str(value_for("nama")).split())
        mark = "" if _is_blank(value_for("mark")) else " ".join(str(value_for("mark")).split()).upper()
        date_value = parse_date(value_for("tanggal"))
        nik = clean_nik(value_for("nik"))
        row = {
            "source_row": row_number,
            "source_values": [
                {"header": str(header), "value": _raw_source_value(values[index] if index < len(values) else None)}
                for index, header in enumerate(source_headers)
                if header is not None
            ],
            "new_employee_id": new_employee_id or None,
            "nik": None if nik is False else nik,
            "cabang": " ".join(str(value_for("cabang") or "").split()),
            "nama": nama,
            "payroll_status": None if _is_blank(value_for("status")) else " ".join(str(value_for("status")).split()).title(),
            "tanggal": None if date_value is False else date_value,
            "mark": mark,
        }
        row_errors = []
        if not nama:
            row_errors.append("Nama kosong.")
        if nik is False:
            row_errors.append("NIK tidak valid.")
        if date_value is False or date_value is None:
            row_errors.append("Tanggal tidak valid.")
        if row_errors:
            conflicts.append({"row": row_number, "new_employee_id": new_employee_id or None, "reason": " ".join(row_errors)})
        else:
            rows.append(row)
    return rows, conflicts
