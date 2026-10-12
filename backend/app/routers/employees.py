# -*- coding: utf-8 -*-
import datetime
from typing import Optional
from fastapi import APIRouter, Body, Depends, HTTPException, Query, File, UploadFile
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from .. import models, schemas, auth
from ..database import get_db
from ..employee_imports import normalize_text, parse_master_workbook, parse_recap_workbook

router = APIRouter(prefix="/employees", tags=["employees"])


from ..config import CONTRACT_REMINDER_DAYS, calculate_tenure, get_today_jakarta


def _name_reference_conflict(db: Session, employee: models.Employee, new_name: str) -> Optional[str]:
    old_key = normalize_text(employee.nama)
    new_key = normalize_text(new_name)
    if old_key == new_key:
        return None

    other_employees = db.query(models.Employee).filter(models.Employee.id != employee.id).all()
    if any(normalize_text(row.nama) in (old_key, new_key) for row in other_employees):
        return "Nama lama/baru dipakai karyawan lain; relasi berbasis nama perlu ditinjau."

    adjustments = db.query(models.EmployeeAdjustment).all()
    if any(normalize_text(row.nama) == new_key for row in adjustments):
        return "Penyesuaian karyawan dengan nama tujuan sudah ada."

    manual_rows = db.query(models.AbsenManual).all()
    old_manual = [row for row in manual_rows if normalize_text(row.nama) == old_key]
    target_periods = {
        (row.tahun, row.bulan)
        for row in manual_rows
        if normalize_text(row.nama) == new_key
    }
    if any((row.tahun, row.bulan) in target_periods for row in old_manual):
        return "Absen manual dengan nama tujuan sudah ada pada periode yang sama."
    return None


def _rename_employee_name_references(db: Session, old_name: str, new_name: str) -> None:
    old_key = normalize_text(old_name)
    if old_key == normalize_text(new_name):
        return
    for row in db.query(models.LeaveRequest).all():
        if normalize_text(row.nama) == old_key:
            row.nama = new_name
    for row in db.query(models.EmployeeAdjustment).all():
        if normalize_text(row.nama) == old_key:
            row.nama = new_name
    for row in db.query(models.AbsenManual).all():
        if normalize_text(row.nama) == old_key:
            row.nama = new_name


def _candidate_dict(candidate: models.EmployeeImportCandidate) -> dict:
    return {
        "id": candidate.id,
        "employee_code": candidate.employee_code,
        "nama": candidate.nama,
        "nik": candidate.nik,
        "join_date": candidate.join_date,
        "uang_makan_override": candidate.uang_makan_override,
        "jabatan": candidate.jabatan,
        "cabang": candidate.cabang,
        "payroll_status": candidate.payroll_status,
        "source_row": candidate.source_row,
        "status": candidate.status,
    }


@router.post("/import/master")
async def import_employee_master(
    file: UploadFile = File(...),
    dry_run: bool = Query(default=True),
    deactivate_missing: bool = Query(default=False),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_hr_master),
):
    """Sinkronkan Master dengan pratinjau aman; karyawan baru menunggu review profil."""
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="File Master harus berformat .xlsx atau .xlsm.")
    try:
        rows, conflicts = parse_master_workbook(await file.read())
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not rows and not conflicts:
        raise HTTPException(status_code=400, detail="Workbook Master tidak berisi data karyawan.")

    employees = db.query(models.Employee).all()
    by_code = {}
    by_nik = {}
    by_name = {}
    for employee in employees:
        if employee.employee_code:
            by_code.setdefault(employee.employee_code, []).append(employee)
        if employee.nik:
            by_nik.setdefault(employee.nik, []).append(employee)
        by_name.setdefault(normalize_text(employee.nama), []).append(employee)

    present_codes = {row["employee_code"] for row in rows}
    present_codes.update(
        conflict["employee_code"]
        for conflict in conflicts
        if conflict.get("employee_code")
    )
    plans = []
    matched_employee_ids = set()
    for row in rows:
        code = row["employee_code"]
        employee = None
        matched_by = None
        code_matches = by_code.get(code, [])
        if len(code_matches) > 1:
            conflicts.append({"row": row["source_row"], "employee_code": code, "reason": "ID JIP duplikat di database."})
            continue
        if code_matches:
            employee, matched_by = code_matches[0], "employee_code"
        elif row["nik"]:
            nik_matches = by_nik.get(row["nik"], [])
            if len(nik_matches) > 1:
                conflicts.append({"row": row["source_row"], "employee_code": code, "reason": "NIK terhubung ke lebih dari satu karyawan di database."})
                continue
            if nik_matches:
                employee, matched_by = nik_matches[0], "nik"
        if employee is None:
            name_matches = by_name.get(normalize_text(row["nama"]), [])
            if len(name_matches) > 1:
                conflicts.append({"row": row["source_row"], "employee_code": code, "reason": "Nama cocok ke lebih dari satu karyawan di database."})
                continue
            if name_matches:
                employee, matched_by = name_matches[0], "nama"

        if employee is None:
            plans.append({"row": row, "employee": None, "matched_by": None})
            continue
        if employee.id in matched_employee_ids:
            conflicts.append({"row": row["source_row"], "employee_code": code, "reason": "Karyawan database sudah dicocokkan dengan baris Master lain."})
            continue
        if employee.employee_code and employee.employee_code != code:
            conflicts.append({"row": row["source_row"], "employee_code": code, "reason": "Karyawan yang cocok sudah memiliki ID JIP berbeda."})
            continue
        if row["nik"]:
            owners = [owner for owner in by_nik.get(row["nik"], []) if owner.id != employee.id]
            if owners:
                conflicts.append({"row": row["source_row"], "employee_code": code, "reason": "NIK sudah dipakai karyawan lain di database."})
                continue
            if matched_by != "employee_code" and employee.nik and employee.nik != row["nik"]:
                conflicts.append({"row": row["source_row"], "employee_code": code, "reason": "NIK berbeda dari karyawan yang cocok; perlu verifikasi manual."})
                continue
        if any(
            other.id != employee.id and normalize_text(other.nama) == normalize_text(row["nama"])
            for other in employees
        ):
            conflicts.append({"row": row["source_row"], "employee_code": code, "reason": "Nama kanonis sudah dipakai karyawan lain."})
            continue
        rename_conflict = _name_reference_conflict(db, employee, row["nama"])
        if rename_conflict:
            conflicts.append({"row": row["source_row"], "employee_code": code, "reason": rename_conflict})
            continue
        plans.append({"row": row, "employee": employee, "matched_by": matched_by})
        matched_employee_ids.add(employee.id)

    candidate_by_code = {
        candidate.employee_code: candidate
        for candidate in db.query(models.EmployeeImportCandidate)
        .all()
    }
    updates = []
    new_candidates = []
    for plan in plans:
        row = plan["row"]
        employee = plan["employee"]
        if employee is None:
            new_candidates.append(row)
            continue
        changed = ["employee_code", "cabang", "active"]
        if normalize_text(employee.nama) != normalize_text(row["nama"]):
            changed.append("nama")
        if row["nik"] and employee.nik != row["nik"]:
            changed.append("nik")
        if row["join_date"] and employee.join_date != row["join_date"]:
            changed.append("join_date")
        if row["uang_makan_override"] is not None and employee.uang_makan_override != row["uang_makan_override"]:
            changed.append("uang_makan_override")
        if row["jabatan"] and employee.jabatan != row["jabatan"]:
            changed.append("jabatan")
        if row["payroll_status"] and employee.payroll_status != row["payroll_status"]:
            changed.append("payroll_status")
        updates.append({
            "row": row,
            "employee_id": employee.id,
            "matched_by": plan["matched_by"],
            "fields": changed,
        })

    removal_candidates = [
        employee for employee in employees
        if employee.employee_code and employee.employee_code not in present_codes
    ]
    deactivation_blocked = bool(deactivate_missing and conflicts)
    source_row_numbers = {
        row["source_row"] for row in rows
    } | {
        conflict["row"] for conflict in conflicts if conflict.get("row") is not None
    }

    if not dry_run:
        try:
            for plan in plans:
                row = plan["row"]
                employee = plan["employee"]
                if employee is None:
                    candidate = candidate_by_code.get(row["employee_code"])
                    if candidate is None:
                        candidate = models.EmployeeImportCandidate(employee_code=row["employee_code"])
                        db.add(candidate)
                        candidate_by_code[row["employee_code"]] = candidate
                    elif candidate.status != "PENDING":
                        candidate.status = "PENDING"
                        candidate.approved_employee_id = None
                    for field in (
                        "nama", "nik", "join_date", "uang_makan_override", "jabatan",
                        "cabang", "payroll_status", "source_row",
                    ):
                        setattr(candidate, field, row[field])
                    continue

                old_name = employee.nama
                if normalize_text(old_name) != normalize_text(row["nama"]):
                    _rename_employee_name_references(db, old_name, row["nama"])
                    employee.nama = row["nama"]
                employee.employee_code = row["employee_code"]
                employee.cabang = row["cabang"]
                employee.active = True
                if row["nik"]:
                    employee.nik = row["nik"]
                if row["join_date"]:
                    employee.join_date = row["join_date"]
                if row["uang_makan_override"] is not None:
                    employee.uang_makan_override = row["uang_makan_override"]
                if row["jabatan"]:
                    employee.jabatan = row["jabatan"]
                if row["payroll_status"]:
                    employee.payroll_status = row["payroll_status"]
                candidate = candidate_by_code.get(row["employee_code"])
                if candidate and candidate.status == "PENDING":
                    candidate.status = "APPROVED"
                    candidate.approved_employee_id = employee.id

            if deactivate_missing and not deactivation_blocked:
                for employee in removal_candidates:
                    employee.active = False
            db.commit()
        except IntegrityError as exc:
            db.rollback()
            raise HTTPException(status_code=409, detail="Sinkronisasi dibatalkan karena konflik ID JIP/NIK yang unik.") from exc
        except Exception:
            db.rollback()
            raise
    else:
        db.rollback()

    return {
        "message": "Pratinjau selesai." if dry_run else "Sinkronisasi selesai; karyawan baru menunggu review HR.",
        "dry_run": dry_run,
        "total_source_rows": len(source_row_numbers),
        "total_updated": len(updates),
        "total_new_pending_review": len(new_candidates),
        "total_conflicts": len(conflicts),
        "total_missing_from_master": len(removal_candidates),
        "deactivate_missing_requested": deactivate_missing,
        "deactivation_blocked_by_conflicts": deactivation_blocked,
        "detail": {
            "updated": updates,
            "new_pending_review": [
                {"source_row": row["source_row"], "employee_code": row["employee_code"], "nama": row["nama"]}
                for row in new_candidates
            ],
            "conflicts": conflicts,
            "missing_from_master": [
                {"employee_id": employee.id, "employee_code": employee.employee_code, "nama": employee.nama}
                for employee in removal_candidates
            ],
        },
    }


@router.get("/import/master/pending")
def list_employee_import_candidates(
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_hr_master),
):
    rows = (
        db.query(models.EmployeeImportCandidate)
        .filter(models.EmployeeImportCandidate.status == "PENDING")
        .order_by(models.EmployeeImportCandidate.employee_code)
        .all()
    )
    return [_candidate_dict(row) for row in rows]


@router.post("/import/master/pending/{candidate_id}/approve", response_model=schemas.EmployeeOut)
def approve_employee_import_candidate(
    candidate_id: int,
    payload: schemas.EmployeeImportApproval,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_hr_master),
):
    candidate = db.query(models.EmployeeImportCandidate).filter(
        models.EmployeeImportCandidate.id == candidate_id,
        models.EmployeeImportCandidate.status == "PENDING",
    ).first()
    if not candidate:
        raise HTTPException(status_code=404, detail="Kandidat impor tidak ditemukan atau sudah diproses.")
    profile = db.query(models.Profile).filter(models.Profile.code == payload.profile_code).first()
    if not profile:
        raise HTTPException(status_code=400, detail=f"Profil '{payload.profile_code}' tidak ditemukan.")

    employment_status = _normalize_employment_status(payload.employment_status)
    if employment_status not in ("PKWTT", "PKWT", "PHL"):
        raise HTTPException(status_code=400, detail="Status kepegawaian harus PKWTT, PKWT, atau PHL.")
    join_date = payload.join_date or candidate.join_date
    if not join_date:
        raise HTTPException(status_code=400, detail="Tanggal Join wajib dilengkapi sebelum kandidat diaktifkan.")
    if candidate.nik and db.query(models.Employee).filter(models.Employee.nik == candidate.nik).first():
        raise HTTPException(status_code=409, detail="NIK kandidat sudah dipakai karyawan lain.")
    if db.query(models.Employee).filter(models.Employee.employee_code == candidate.employee_code).first():
        raise HTTPException(status_code=409, detail="ID JIP kandidat sudah terdaftar sebagai karyawan.")

    contract = None
    if employment_status == "PKWT":
        if not payload.contract_start_date or not payload.contract_end_date:
            raise HTTPException(status_code=400, detail="Tanggal mulai dan akhir kontrak wajib untuk PKWT.")
        if payload.contract_end_date <= payload.contract_start_date:
            raise HTTPException(status_code=400, detail="Tanggal akhir kontrak harus setelah tanggal mulai.")
        contract = (payload.contract_start_date, payload.contract_end_date)

    employee = models.Employee(
        employee_code=candidate.employee_code,
        nama=candidate.nama,
        nik=candidate.nik,
        profile_code=profile.code,
        cabang=candidate.cabang,
        jabatan=candidate.jabatan,
        payroll_status=candidate.payroll_status,
        uang_makan_override=candidate.uang_makan_override,
        join_date=join_date,
        employment_status=employment_status,
        active=True,
        profile_needs_review=False,
    )
    db.add(employee)
    db.flush()
    if contract:
        db.add(models.EmploymentContract(
            employee_id=employee.id,
            contract_number=1,
            start_date=contract[0],
            end_date=contract[1],
            status="ACTIVE",
            keterangan=payload.contract_keterangan or "Kontrak pertama (impor Master)",
        ))
    candidate.status = "APPROVED"
    candidate.approved_employee_id = employee.id
    db.commit()
    db.refresh(employee)
    _enrich_employee_data(employee)
    return employee


@router.post("/import/recap")
async def import_clean_attendance_recap(
    file: UploadFile = File(...),
    dry_run: bool = Query(default=True),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_hr_master),
):
    """Impor kode cuti/sakit terverifikasi; T tidak menjadi hitungan scan/keterlambatan."""
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="File Rekap harus berformat .xlsx atau .xlsm.")
    try:
        rows, conflicts = parse_recap_workbook(await file.read())
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    employee_by_code = {
        employee.employee_code: employee
        for employee in db.query(models.Employee).filter(models.Employee.employee_code.isnot(None)).all()
    }
    event_categories = {
        "C": ("CUTI_TAHUNAN", 1.0),
        "CS": ("CUTI_SETENGAH_HARI", 0.5),
        "S": ("SAKIT", 1.0),
    }
    existing_import_keys = {
        request.import_key
        for request in db.query(models.LeaveRequest).filter(models.LeaveRequest.import_key.isnot(None)).all()
    }
    all_leaves = db.query(models.LeaveRequest).all()
    seen_keys = set()
    prepared = []
    ignored_late = 0
    duplicates = 0

    for row in rows:
        code = row["employee_code"]
        employee = employee_by_code.get(code)
        if not employee:
            conflicts.append({"row": row["source_row"], "employee_code": code, "reason": "ID JIP belum terdaftar di database; sinkronkan Master dahulu."})
            continue
        if row["nik"] and employee.nik and row["nik"] != employee.nik:
            conflicts.append({"row": row["source_row"], "employee_code": code, "reason": "NIK Rekap berbeda dari Master/database."})
            continue
        if normalize_text(row["nama"]) != normalize_text(employee.nama):
            conflicts.append({"row": row["source_row"], "employee_code": code, "reason": "Nama Rekap berbeda dari Master/database."})
            continue
        if row["mark"] == "T":
            ignored_late += 1
            continue
        if row["mark"] not in event_categories:
            conflicts.append({"row": row["source_row"], "employee_code": code, "reason": f"Kode keterangan tidak dikenal: {row['mark'] or '(kosong)'}."})
            continue

        category, days = event_categories[row["mark"]]
        import_key = f"rekap-bersih:{code}:{row['tanggal'].isoformat()}:{category}"
        if import_key in seen_keys or import_key in existing_import_keys:
            duplicates += 1
            continue
        already_recorded = any(
            normalize_text(leave.nama) == normalize_text(employee.nama)
            and leave.tanggal_mulai == row["tanggal"]
            and leave.tanggal_selesai == row["tanggal"]
            and leave.kategori == category
            for leave in all_leaves
        )
        if already_recorded:
            duplicates += 1
            seen_keys.add(import_key)
            continue
        seen_keys.add(import_key)
        prepared.append({
            "source_row": row["source_row"],
            "employee": employee,
            "employee_code": code,
            "tanggal": row["tanggal"],
            "mark": row["mark"],
            "kategori": category,
            "jumlah_hari": days,
            "import_key": import_key,
        })

    source_row_numbers = {
        row["source_row"] for row in rows
    } | {
        conflict["row"] for conflict in conflicts if conflict.get("row") is not None
    }

    if not dry_run:
        from ..leave_logic import record_leave_used
        try:
            for item in prepared:
                leave = models.LeaveRequest(
                    import_key=item["import_key"],
                    nama=item["employee"].nama,
                    kategori=item["kategori"],
                    tanggal_mulai=item["tanggal"],
                    tanggal_selesai=item["tanggal"],
                    jumlah_hari=item["jumlah_hari"],
                    alasan=f"Diimpor dari Rekap Kehadiran Bersih ({item['mark']}).",
                    status="APPROVED",
                    catatan_hr="Histori impor dari Rekap Kehadiran Bersih.",
                    approved_by="Import Rekap Kehadiran Bersih",
                    approved_at=datetime.datetime.utcnow(),
                )
                db.add(leave)
                db.flush()
                if item["kategori"] in ("CUTI_TAHUNAN", "CUTI_SETENGAH_HARI"):
                    record_leave_used(
                        db,
                        item["employee"].id,
                        item["tanggal"].year,
                        item["jumlah_hari"],
                        leave.id,
                        f"{item['kategori']} {item['tanggal']} (import Rekap)",
                        commit=False,
                    )
            db.commit()
        except IntegrityError as exc:
            db.rollback()
            raise HTTPException(status_code=409, detail="Impor dibatalkan karena ada baris Rekap yang sudah pernah diimpor.") from exc
        except Exception:
            db.rollback()
            raise
    else:
        db.rollback()

    return {
        "message": "Pratinjau selesai." if dry_run else "Impor Rekap selesai.",
        "dry_run": dry_run,
        "total_source_rows": len(source_row_numbers),
        "total_to_import": len(prepared),
        "total_duplicate_or_existing": duplicates,
        "total_t_ignored": ignored_late,
        "total_conflicts": len(conflicts),
        "detail": {
            "to_import": [
                {
                    "source_row": item["source_row"],
                    "employee_code": item["employee_code"],
                    "tanggal": item["tanggal"],
                    "kategori": item["kategori"],
                    "jumlah_hari": item["jumlah_hari"],
                }
                for item in prepared
            ],
            "conflicts": conflicts,
        },
    }


@router.get("/whatsapp-identities")
def list_whatsapp_identities(
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_hr_master),
):
    """Daftar linking nomor WhatsApp untuk ditinjau HR."""
    rows = db.query(models.WhatsAppIdentity).order_by(models.WhatsAppIdentity.updated_at.desc()).all()
    return [{
        "id": row.id,
        "phone_number": row.phone_number,
        "employee_id": row.employee_id,
        "employee_name": row.employee.nama if row.employee else None,
        "employee_nik_last4": row.employee.nik[-4:] if row.employee and row.employee.nik else None,
        "status": row.status,
        "verified_at": row.verified_at,
        "last_seen_at": row.last_seen_at,
        "failed_attempts": row.failed_attempts,
    } for row in rows]


@router.patch("/whatsapp-identities/{identity_id}/revoke")
def revoke_whatsapp_identity(
    identity_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.require_hr_master),
):
    """Cabut linking WhatsApp; nomor dapat login ulang setelah proses verifikasi."""
    identity = db.query(models.WhatsAppIdentity).filter(models.WhatsAppIdentity.id == identity_id).first()
    if not identity:
        raise HTTPException(status_code=404, detail="Linking WhatsApp tidak ditemukan.")
    identity.status = "REVOKED"
    db.add(models.WhatsAppAuditLog(
        phone_number=identity.phone_number,
        employee_id=identity.employee_id,
        event_type="HR_REVOKE",
        detail=f"Linking dicabut oleh {current_user.nama or 'HR Master'}",
    ))
    db.commit()
    return {"status": "sukses", "identity_id": identity.id, "new_status": identity.status}


@router.get("/whatsapp-audit")
def list_whatsapp_audit(
    limit: int = Query(default=100, ge=1, le=500),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_hr_master),
):
    """Histori login, logout, kegagalan, dan perubahan linking WhatsApp."""
    rows = (
        db.query(models.WhatsAppAuditLog)
        .order_by(models.WhatsAppAuditLog.created_at.desc())
        .limit(limit)
        .all()
    )
    return [{
        "id": row.id,
        "phone_number": row.phone_number,
        "employee_id": row.employee_id,
        "employee_name": row.employee.nama if row.employee else None,
        "event_type": row.event_type,
        "detail": row.detail,
        "created_at": row.created_at,
    } for row in rows]


@router.patch("/whatsapp-identities/{identity_id}/move")
def move_whatsapp_identity(
    identity_id: int,
    payload: dict = Body(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.require_hr_master),
):
    """Pindahkan linking ke karyawan lain setelah verifikasi manual oleh HR."""
    target_employee_code = str(payload.get("employee_code") or "").strip()
    reason = str(payload.get("reason") or "").strip()
    if not target_employee_code or not reason:
        raise HTTPException(status_code=400, detail="employee_code dan alasan pemindahan wajib diisi.")

    identity = db.query(models.WhatsAppIdentity).filter(models.WhatsAppIdentity.id == identity_id).first()
    target = db.query(models.Employee).filter(
        models.Employee.employee_code == target_employee_code,
    ).first()
    if not identity:
        raise HTTPException(status_code=404, detail="Linking WhatsApp tidak ditemukan.")
    if not target:
        raise HTTPException(status_code=404, detail="Kode Karyawan tujuan tidak ditemukan.")

    conflict = db.query(models.WhatsAppIdentity).filter(
        models.WhatsAppIdentity.employee_id == target.id,
        models.WhatsAppIdentity.status == "ACTIVE",
        models.WhatsAppIdentity.id != identity.id,
    ).first()
    if conflict:
        raise HTTPException(status_code=409, detail="Karyawan tujuan sudah memiliki linking WhatsApp aktif.")

    previous_employee_id = identity.employee_id
    identity.employee_id = target.id
    identity.status = "ACTIVE"
    identity.verified_at = datetime.datetime.utcnow()
    identity.failed_attempts = 0
    identity.locked_until = None
    db.add(models.WhatsAppAuditLog(
        phone_number=identity.phone_number,
        employee_id=target.id,
        event_type="HR_MOVE",
        detail=(
            f"Linking dipindahkan dari employee_id={previous_employee_id} "
            f"ke employee_id={target.id} oleh {current_user.nama or 'HR Master'}: {reason}"
        ),
    ))
    db.commit()
    return {
        "status": "sukses",
        "identity_id": identity.id,
        "employee_id": target.id,
        "employee_code": target.employee_code,
        "new_status": identity.status,
    }


def _normalize_employment_status(value: Optional[str]) -> str:
    """Normalisasi status kepegawaian agar label baru aman untuk data lama."""
    normalized = (value or "PKWTT").upper().strip()
    aliases = {
        "TETAP": "PKWTT",
        "KARYAWAN_TETAP": "PKWTT",
        "PKWTT": "PKWTT",
        "PKWT": "PKWT",
        "PHL": "PHL",
    }
    return aliases.get(normalized, normalized)


def _enrich_employee_data(emp: models.Employee) -> models.Employee:
    """Helper untuk menghitung tenure & contract_reminder_status pada objek Employee."""
    emp.employment_status = _normalize_employment_status(emp.employment_status)
    emp.tenure_display = calculate_tenure(emp.join_date)["display"]

    today = get_today_jakarta()
    if emp.employment_status == "PKWT" and emp.contracts:
        active_contract = next((c for c in emp.contracts if c.status == "ACTIVE"), emp.contracts[0] if emp.contracts else None)
        if active_contract:
            days_left = (active_contract.end_date - today).days
            if days_left < 0:
                status_label = "EXPIRED"
                badge_color = "red"
                msg = f"Kontrak #{active_contract.contract_number} telah berakhir {abs(days_left)} hari lalu ({active_contract.end_date.strftime('%d-%m-%Y')})"
            elif days_left <= CONTRACT_REMINDER_DAYS:
                status_label = "EXPIRING_SOON"
                badge_color = "amber"
                msg = f"Sisa {days_left} hari! Kontrak #{active_contract.contract_number} berakhir {active_contract.end_date.strftime('%d-%m-%Y')}"
            else:
                status_label = "ACTIVE"
                badge_color = "green"
                msg = f"Kontrak #{active_contract.contract_number} Aktif ({days_left} hari tersisa)"

            emp.contract_reminder_status = {
                "contract_id": active_contract.id,
                "contract_number": active_contract.contract_number,
                "start_date": active_contract.start_date.isoformat(),
                "end_date": active_contract.end_date.isoformat(),
                "days_left": days_left,
                "status_label": status_label,
                "badge_color": badge_color,
                "message": msg,
                "is_expired": days_left < 0,
                "is_expiring": 0 <= days_left <= CONTRACT_REMINDER_DAYS,
            }
        else:
            emp.contract_reminder_status = None
    else:
        emp.contract_reminder_status = None

    return emp


@router.get("", response_model=list[schemas.EmployeeOut])
def list_employees(
    profile_code: Optional[str] = Query(default=None),
    employment_status: Optional[str] = Query(default=None),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_user),
):
    q = db.query(models.Employee)
    if profile_code:
        q = q.filter(models.Employee.profile_code == profile_code)
    if employment_status:
        q = q.filter(models.Employee.employment_status == _normalize_employment_status(employment_status))

    employees = q.order_by(models.Employee.nama).all()
    for emp in employees:
        _enrich_employee_data(emp)
    return employees


@router.get("/contracts/expiring")
def get_expiring_contracts(
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_user),
):
    """
    Mengambil daftar karyawan PKWT yang kontraknya berakhir dalam H-10 (atau sudah lewat)
    dan belum diperpanjang / belum diangkat menjadi TETAP.
    Diurutkan dari yang paling dekat tanggal berakhirnya.
    """
    pkwt_emps = (
        db.query(models.Employee)
        .filter(models.Employee.employment_status == "PKWT", models.Employee.active == True)
        .all()
    )

    expiring_list = []
    for emp in pkwt_emps:
        _enrich_employee_data(emp)
        rem = emp.contract_reminder_status
        if rem and (rem["is_expiring"] or rem["is_expired"]):
            expiring_list.append({
                "employee_id": emp.id,
                "nama": emp.nama,
                "cabang": emp.cabang,
                "profile_code": emp.profile_code,
                "join_date": emp.join_date,
                "tenure_display": emp.tenure_display,
                "contract_id": rem["contract_id"],
                "contract_number": rem["contract_number"],
                "start_date": rem["start_date"],
                "end_date": rem["end_date"],
                "days_left": rem["days_left"],
                "status_label": rem["status_label"],
                "is_expired": rem["is_expired"],
                "message": rem["message"],
            })

    expiring_list.sort(key=lambda x: x["days_left"])
    return {
        "reminder_days_config": CONTRACT_REMINDER_DAYS,
        "total": len(expiring_list),
        "data": expiring_list,
    }


# ─────────────────────────────────────────────────────────────
# Employee Adjustments (Bonus Lain / Potongan Lain)
# ─────────────────────────────────────────────────────────────

@router.get("/adjustments", response_model=list[schemas.EmployeeAdjustmentOut])
def list_adjustments(
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_user),
):
    """Mendapatkan semua data penyesuaian (bonus/potongan lain-lain) karyawan."""
    return db.query(models.EmployeeAdjustment).order_by(models.EmployeeAdjustment.nama).all()


@router.patch("/adjustments", response_model=schemas.EmployeeAdjustmentOut)
def upsert_adjustment(
    payload: schemas.EmployeeAdjustmentUpdate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_user),
):
    """Membuat atau memperbarui data penyesuaian karyawan (upsert by nama)."""
    nama_key = payload.nama.strip()
    if not nama_key:
        raise HTTPException(status_code=400, detail="Nama karyawan wajib diisi.")

    adj = db.query(models.EmployeeAdjustment).filter(
        models.EmployeeAdjustment.nama == nama_key
    ).first()

    if adj is None:
        adj = models.EmployeeAdjustment(
            nama=nama_key,
            bonus_lain=payload.bonus_lain if payload.bonus_lain is not None else 0.0,
            potongan_lain=payload.potongan_lain if payload.potongan_lain is not None else 0.0,
            catatan=payload.catatan,
        )
        db.add(adj)
    else:
        if payload.bonus_lain is not None:
            adj.bonus_lain = payload.bonus_lain
        if payload.potongan_lain is not None:
            adj.potongan_lain = payload.potongan_lain
        if payload.catatan is not None:
            adj.catatan = payload.catatan

    db.commit()
    db.refresh(adj)
    return adj


# ─────────────────────────────────────────────────────────────
# Absen Manual (Hari Kerja Manual per Periode)
# ─────────────────────────────────────────────────────────────

@router.put("/absen-manual", response_model=schemas.AbsenManualOut)
def upsert_absen_manual(
    payload: schemas.AbsenManualUpdate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_user),
):
    """Membuat atau memperbarui jumlah hari absen manual karyawan untuk periode tertentu (upsert)."""
    nama_key = payload.nama.strip()
    if not nama_key:
        raise HTTPException(status_code=400, detail="Nama karyawan wajib diisi.")

    record = db.query(models.AbsenManual).filter(
        models.AbsenManual.nama == nama_key,
        models.AbsenManual.tahun == payload.tahun,
        models.AbsenManual.bulan == payload.bulan,
    ).first()

    if record is None:
        record = models.AbsenManual(
            nama=nama_key,
            tahun=payload.tahun,
            bulan=payload.bulan,
            jumlah=payload.jumlah,
        )
        db.add(record)
    else:
        record.jumlah = payload.jumlah

    db.commit()
    db.refresh(record)
    return record


# ─────────────────────────────────────────────────────────────
# Template Karyawan
# ─────────────────────────────────────────────────────────────

@router.get("/template")
def download_employee_template(
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_hr_master),
):
    """Download template master karyawan yang mengikuti field HR aplikasi."""
    import io
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from fastapi.responses import Response

    profiles = db.query(models.Profile).order_by(models.Profile.code).all()
    branches = sorted({
        branch.strip()
        for profile in profiles
        for branch in (profile.cabang or [])
        if isinstance(branch, str) and branch.strip()
    } | {
        employee.cabang.strip()
        for employee in db.query(models.Employee.cabang).all()
        if employee.cabang and employee.cabang.strip()
    })

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data Karyawan"
    headers = [
        "Kode Karyawan", "Nama Karyawan", "NIK", "Cabang", "Kode Profil",
        "Status Kepegawaian", "Tanggal Masuk", "Status Aktif",
        "Jabatan", "ID Mesin Absensi",
    ]
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:J1"
    widths = [20, 30, 20, 24, 22, 24, 18, 16, 24, 20]
    for index, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = sheet.cell(row=1, column=index)
        cell.fill = PatternFill("solid", fgColor="284B3A")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width

    for row in range(2, 502):
        sheet.cell(row=row, column=3).number_format = "@"
        sheet.cell(row=row, column=7).number_format = "yyyy-mm-dd"

    status_validation = DataValidation(type="list", formula1='"PKWTT,PKWT,PHL"', allow_blank=False)
    active_validation = DataValidation(type="list", formula1='"Aktif,Nonaktif"', allow_blank=True)
    sheet.add_data_validation(status_validation)
    sheet.add_data_validation(active_validation)
    status_validation.add("F2:F501")
    active_validation.add("H2:H501")

    reference = workbook.create_sheet("Referensi")
    reference.append(["Kode Profil", "Nama Profil"])
    for profile in profiles:
        reference.append([profile.code, profile.nama])
    reference.append([])
    reference.append(["Status Kepegawaian", "Arti"])
    reference.append(["PKWTT", "Perjanjian Kerja Waktu Tidak Tertentu"])
    reference.append(["PKWT", "Perjanjian Kerja Waktu Tertentu"])
    reference.append(["PHL", "Pekerja Harian Lepas"])
    reference.append([])
    reference.append(["Status Aktif", "Arti"])
    reference.append(["Aktif", "Karyawan aktif"])
    reference.append(["Nonaktif", "Karyawan nonaktif"])
    reference.append([])
    reference.append(["Cabang yang sudah dikenal aplikasi"])
    for branch in branches:
        reference.append([branch])
    reference.freeze_panes = "A2"
    reference.column_dimensions["A"].width = 32
    reference.column_dimensions["B"].width = 55

    guide = workbook.create_sheet("Petunjuk")
    instructions = [
        ["PANDUAN IMPOR MASTER KARYAWAN"],
        ["Isi satu baris per karyawan pada sheet Data Karyawan."],
        ["Kode Karyawan, Nama Karyawan, Cabang, Kode Profil, dan Status Kepegawaian wajib diisi."],
        ["Gunakan Kode Profil persis seperti yang tercantum pada sheet Referensi."],
        ["Status Kepegawaian hanya PKWTT, PKWT, atau PHL; Status Aktif hanya Aktif atau Nonaktif."],
        ["NIK boleh kosong; jika diisi harus 16 digit dan tidak boleh digunakan karyawan lain."],
        ["Tanggal Masuk wajib untuk karyawan baru. Format yang disarankan: YYYY-MM-DD."],
        ["Gaji, BPJS, dan tanggal akhir kontrak tidak dicatat melalui template ini."],
        ["PKWT tanpa kontrak aktif akan diimpor sebagai nonaktif. Lengkapi tanggal kontrak melalui form Karyawan sebelum mengaktifkannya."],
        ["File lama dengan format Status/Nama/Uang Makan tidak lagi didukung."],
    ]
    for row in instructions:
        guide.append(row)
    guide.column_dimensions["A"].width = 120
    guide["A1"].font = Font(bold=True, size=14, color="284B3A")
    for cell in guide["A"]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    guide.freeze_panes = "A2"

    output = io.BytesIO()
    workbook.save(output)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="template_karyawan.xlsx"'},
    )


# ─────────────────────────────────────────────────────────────
# Template & Import NIK / Join Date / Waktu Berakhir Kontrak
# ─────────────────────────────────────────────────────────────

@router.get("/template/nik-import")
def download_nik_import_template(
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_hr_master),
):
    """Download workbook untuk memperbarui data master dengan Kode Karyawan sebagai kunci."""
    import io
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    from fastapi.responses import Response

    employees = (
        db.query(models.Employee)
        .order_by(models.Employee.cabang, models.Employee.nama)
        .all()
    )

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Pembaruan Master"
    headers = [
        "Kode Karyawan", "Nama Karyawan", "NIK", "Cabang", "Kode Profil",
        "Status Kepegawaian", "Tanggal Masuk", "Status Aktif", "Jabatan",
        "ID Mesin Absensi",
    ]
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:J{max(1, len(employees) + 1)}"
    widths = [20, 30, 20, 24, 22, 24, 18, 16, 24, 20]
    for index, (width, _) in enumerate(zip(widths, headers), start=1):
        cell = sheet.cell(row=1, column=index)
        cell.fill = PatternFill("solid", fgColor="284B3A")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width

    for row_index, employee in enumerate(employees, start=2):
        values = [
            employee.employee_code or "",
            employee.nama,
            employee.nik or "",
            employee.cabang or "",
            employee.profile_code or "",
            _normalize_employment_status(employee.employment_status),
            employee.join_date.isoformat() if employee.join_date else "",
            "Aktif" if employee.active else "Nonaktif",
            employee.jabatan or "",
            employee.id_mesin or "",
        ]
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.fill = PatternFill(
                "solid",
                fgColor="D9E2F3" if column_index <= 2 else "FFFDE7",
            )
            cell.alignment = Alignment(vertical="center")
        sheet.cell(row=row_index, column=3).number_format = "@"
        sheet.cell(row=row_index, column=7).number_format = "yyyy-mm-dd"

    status_validation = DataValidation(type="list", formula1='"PKWTT,PKWT,PHL"', allow_blank=True)
    active_validation = DataValidation(type="list", formula1='"Aktif,Nonaktif"', allow_blank=True)
    sheet.add_data_validation(status_validation)
    sheet.add_data_validation(active_validation)
    status_validation.add(f"F2:F{max(501, len(employees) + 1)}")
    active_validation.add(f"H2:H{max(501, len(employees) + 1)}")

    guide = workbook.create_sheet("Petunjuk")
    for line in [
        ["PETUNJUK PEMBARUAN MASTER KARYAWAN"],
        ["Setiap baris dicocokkan berdasarkan Kode Karyawan; jangan mengubah kolom Kode Karyawan dan Nama Karyawan."],
        ["Kolom kosong selain kunci/nama tidak akan menghapus nilai yang sudah tersimpan."],
        ["NIK harus 16 digit; Tanggal Masuk menggunakan format YYYY-MM-DD."],
        ["Kode Profil harus sama persis dengan kode pada menu Profil/Divisi."],
        ["Status Kepegawaian: PKWTT, PKWT, atau PHL. Status Aktif: Aktif atau Nonaktif."],
        ["Template ini tidak mengubah tanggal kontrak. Karyawan PKWT tanpa kontrak aktif tetap nonaktif sampai kontrak dilengkapi di form Karyawan."],
        ["Format template lama yang mencocokkan berdasarkan nama tidak lagi didukung."],
    ]:
        guide.append(line)
    guide.column_dimensions["A"].width = 120
    guide["A1"].font = Font(bold=True, size=14, color="284B3A")
    for cell in guide["A"]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    guide.freeze_panes = "A2"

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="template_update_master_karyawan.xlsx"'},
    )


@router.post("/import/nik")
async def import_nik_from_excel(
    file: UploadFile = File(...),
    dry_run: bool = Query(default=False),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_hr_master),
):
    """Update master karyawan melalui workbook terbaru dengan Kode Karyawan sebagai kunci."""
    import io
    import openpyxl
    from ..employee_imports import clean_nik, normalize_header, parse_date

    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="File pembaruan master harus berformat .xlsx atau .xlsm.")
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(await file.read()), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="File Excel tidak valid atau rusak.") from exc

    if "Pembaruan Master" not in workbook.sheetnames:
        raise HTTPException(status_code=400, detail="Workbook harus memakai sheet 'Pembaruan Master' dari template terbaru.")
    sheet = workbook["Pembaruan Master"]
    headers = {
        normalize_header(cell.value): index
        for index, cell in enumerate(sheet[1])
        if cell.value is not None
    }
    required_headers = {
        "kode karyawan", "nama karyawan", "nik", "cabang", "kode profil",
        "status kepegawaian", "tanggal masuk", "status aktif",
    }
    if not required_headers.issubset(headers):
        raise HTTPException(status_code=400, detail="Kolom workbook tidak sesuai template pembaruan master terbaru.")

    profiles = {
        profile.code.strip().casefold(): profile.code
        for profile in db.query(models.Profile).all()
    }

    results = {
        "updated": [],
        "skipped": [],
        "not_found": [],
        "nik_conflict": [],
        "invalid_nik": [],
        "invalid_status": [],
        "invalid_profile": [],
        "invalid_join_date": [],
        "contract_required": [],
    }

    def value_for(values, name):
        index = headers[name]
        return values[index] if index < len(values) else None

    rows = 0
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value is not None for value in values):
            continue
        code = str(value_for(values, "kode karyawan") or "").strip()
        if not code:
            results["skipped"].append({"row": row_number, "reason": "Kode Karyawan kosong."})
            continue
        rows += 1
        emp = db.query(models.Employee).filter(models.Employee.employee_code == code).first()
        if not emp:
            results["not_found"].append(code)
            continue

        pending_updates = {}
        updated_fields = []

        nik_raw = value_for(values, "nik")
        if nik_raw not in (None, ""):
            nik = clean_nik(nik_raw)
            if not nik:
                results["invalid_nik"].append({"employee_code": code, "nik_raw": str(nik_raw)})
            else:
                conflict = db.query(models.Employee).filter(
                    models.Employee.nik == nik,
                    models.Employee.id != emp.id,
                ).first()
                if conflict:
                    results["nik_conflict"].append({
                        "employee_code": code,
                        "nik": nik,
                        "konflik_dengan": conflict.nama,
                    })
                else:
                    pending_updates["nik"] = nik
                    updated_fields.append("NIK")

        join_raw = value_for(values, "tanggal masuk")
        if join_raw not in (None, ""):
            join_date = parse_date(join_raw)
            if join_date is False:
                results["invalid_join_date"].append({"employee_code": code, "value": str(join_raw)})
            else:
                pending_updates["join_date"] = join_date
                updated_fields.append("Tanggal Masuk")

        branch = str(value_for(values, "cabang") or "").strip()
        if branch:
            pending_updates["cabang"] = branch
            updated_fields.append("Cabang")

        profile_raw = str(value_for(values, "kode profil") or "").strip()
        if profile_raw:
            profile_code = profiles.get(profile_raw.casefold())
            if profile_code:
                pending_updates["profile_code"] = profile_code
                updated_fields.append("Profil")
            else:
                results["invalid_profile"].append({"employee_code": code, "profile_code": profile_raw})

        status_raw = str(value_for(values, "status kepegawaian") or "").strip()
        if status_raw:
            status = _normalize_employment_status(status_raw)
            if status not in ("PKWTT", "PKWT", "PHL"):
                results["invalid_status"].append({"employee_code": code, "status": status_raw})
            else:
                pending_updates["employment_status"] = status
                updated_fields.append("Status Kepegawaian")

        active_raw = str(value_for(values, "status aktif") or "").strip().casefold()
        if active_raw:
            active_values = {
                "aktif": True, "active": True, "ya": True, "true": True, "1": True,
                "nonaktif": False, "non-aktif": False, "inactive": False, "tidak": False, "false": False, "0": False,
            }
            if active_raw not in active_values:
                results["skipped"].append({"row": row_number, "employee_code": code, "reason": "Status Aktif tidak valid."})
                continue
            pending_updates["active"] = active_values[active_raw]
            updated_fields.append("Status Aktif")

        for field, header in (("jabatan", "jabatan"), ("id_mesin", "id mesin absensi")):
            raw_value = value_for(values, header) if header in headers else None
            if raw_value not in (None, ""):
                pending_updates[field] = str(raw_value).strip()
                updated_fields.append("Jabatan" if field == "jabatan" else "ID Mesin Absensi")

        effective_status = pending_updates.get("employment_status", _normalize_employment_status(emp.employment_status))
        active_contract = next((contract for contract in emp.contracts if contract.status == "ACTIVE"), None)
        if effective_status == "PKWT" and not active_contract:
            pending_updates["active"] = False
            if "Status Aktif" not in updated_fields:
                updated_fields.append("Status Aktif")
            results["contract_required"].append(code)

        if updated_fields:
            if not dry_run:
                for field, value in pending_updates.items():
                    setattr(emp, field, value)
            results["updated"].append({"nama": emp.nama, "employee_code": code, "fields": ", ".join(updated_fields)})
        else:
            results["skipped"].append({"row": row_number, "employee_code": code, "reason": "Tidak ada perubahan."})

    if dry_run:
        db.rollback()
    else:
        db.commit()

    return {
        "message": "Pratinjau selesai." if dry_run else "Pembaruan master selesai.",
        "dry_run": dry_run,
        "total_updated": len(results["updated"]),
        "total_skipped": len(results["skipped"]),
        "total_not_found": len(results["not_found"]),
        "total_nik_conflict": len(results["nik_conflict"]),
        "total_invalid_nik": len(results["invalid_nik"]),
        "total_invalid_status": len(results["invalid_status"]),
        "total_invalid_profile": len(results["invalid_profile"]),
        "total_invalid_join_date": len(results["invalid_join_date"]),
        "total_contract_required": len(results["contract_required"]),
        "total_unmapped_branch": 0,
        "total_source_rows": rows,
        "detail": results,
    }


@router.get("/{employee_id}", response_model=schemas.EmployeeOut)
def get_employee_detail(
    employee_id: int,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_user),
):
    emp = db.query(models.Employee).filter(models.Employee.id == employee_id).first()
    if not emp:
        raise HTTPException(status_code=404, detail="Karyawan tidak ditemukan.")
    _enrich_employee_data(emp)
    return emp


@router.post("", response_model=schemas.EmployeeOut)
def create_employee(
    payload: schemas.EmployeeCreate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_hr_master),
):
    """HANYA HR Master yang boleh menambahkan karyawan baru ke suatu profil."""
    profile = db.query(models.Profile).filter(models.Profile.code == payload.profile_code).first()
    if not profile:
        raise HTTPException(status_code=400, detail=f"Profil '{payload.profile_code}' tidak ditemukan.")

    emp_status = _normalize_employment_status(payload.employment_status)
    if emp_status not in ("PKWTT", "PKWT", "PHL"):
        raise HTTPException(status_code=400, detail="Status kepegawaian harus 'PKWTT', 'PKWT', atau 'PHL'.")

    # Validasi join_date: Wajib untuk semua karyawan baru
    if not payload.join_date:
        raise HTTPException(status_code=400, detail="Tanggal masuk / bergabung (join_date) wajib diisi untuk semua karyawan.")

    # Validasi & Handling Kontrak untuk karyawan PKWT
    contract_data = None
    if emp_status == "PKWT":
        if not payload.contract_start_date or not payload.contract_end_date:
            raise HTTPException(
                status_code=400,
                detail="Untuk karyawan PKWT, Tanggal Mulai Kontrak & Tanggal Selesai Kontrak wajib diisi."
            )
        if payload.contract_end_date <= payload.contract_start_date:
            raise HTTPException(
                status_code=400,
                detail="Tanggal selesai kontrak (contract_end_date) harus setelah tanggal mulai kontrak (contract_start_date)."
            )
        contract_data = {
            "start_date": payload.contract_start_date,
            "end_date": payload.contract_end_date,
            "keterangan": payload.contract_keterangan or "Kontrak Pertama (PKWT-1)",
        }

    emp_dict = payload.dict(exclude={
        "contract_start_date",
        "contract_end_date",
        "contract_keterangan",
        "phl_period_start_date",
    })
    emp_dict["employment_status"] = emp_status
    emp = models.Employee(**emp_dict)
    db.add(emp)
    db.commit()
    db.refresh(emp)

    if contract_data:
        contract = models.EmploymentContract(
            employee_id=emp.id,
            contract_number=1,
            start_date=contract_data["start_date"],
            end_date=contract_data["end_date"],
            status="ACTIVE",
            keterangan=contract_data["keterangan"],
        )
        db.add(contract)
        db.commit()
        db.refresh(emp)

    _enrich_employee_data(emp)
    return emp


@router.put("/{employee_id}", response_model=schemas.EmployeeOut)
def update_employee(
    employee_id: int,
    payload: schemas.EmployeeUpdate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_hr_master),
):
    emp = db.query(models.Employee).filter(models.Employee.id == employee_id).first()
    if not emp:
        raise HTTPException(status_code=404, detail="Karyawan tidak ditemukan.")

    profile = db.query(models.Profile).filter(models.Profile.code == payload.profile_code).first()
    if not profile:
        raise HTTPException(status_code=400, detail=f"Profil '{payload.profile_code}' tidak ditemukan.")

    emp.employment_status = _normalize_employment_status(emp.employment_status)
    new_status = _normalize_employment_status(payload.employment_status)
    if new_status not in ("PKWTT", "PKWT", "PHL"):
        raise HTTPException(status_code=400, detail="Status kepegawaian harus 'PKWTT', 'PKWT', atau 'PHL'.")

    is_phl_to_pkwt = emp.employment_status == "PHL" and new_status == "PKWT"
    phl_period_start = None
    phl_period_end = None
    if is_phl_to_pkwt:
        if not payload.contract_start_date or not payload.contract_end_date:
            raise HTTPException(
                status_code=400,
                detail="Perubahan PHL ke PKWT wajib memiliki tanggal mulai dan selesai kontrak.",
            )
        if payload.contract_end_date <= payload.contract_start_date:
            raise HTTPException(
                status_code=400,
                detail="Tanggal selesai kontrak harus setelah tanggal mulai kontrak.",
            )
        phl_period_start = payload.phl_period_start_date or emp.join_date
        if not phl_period_start:
            raise HTTPException(
                status_code=400,
                detail="Tanggal mulai periode PHL wajib diisi untuk mencatat transisi ke PKWT.",
            )
        phl_period_end = payload.contract_start_date - datetime.timedelta(days=1)
        if phl_period_start > phl_period_end:
            raise HTTPException(
                status_code=400,
                detail="Periode PHL harus berakhir sebelum tanggal mulai kontrak PKWT.",
            )
        if any(contract.status == "ACTIVE" for contract in emp.contracts):
            raise HTTPException(
                status_code=409,
                detail="Karyawan PHL sudah memiliki kontrak aktif; periksa histori kontrak sebelum mengubah status.",
            )

    # Jika diubah dari PKWT ke PKWTT (Pengangkatan Karyawan Tetap)
    if emp.employment_status == "PKWT" and new_status == "PKWTT":
        for c in emp.contracts:
            if c.status == "ACTIVE":
                c.status = "PROMOTED_TO_PERMANENT"
                c.keterangan = (c.keterangan or "") + " [Diangkat menjadi PKWTT]"

    # Jika status PKWT & ada input tanggal kontrak baru/update
    if new_status == "PKWT":
        if is_phl_to_pkwt:
            contract_number = max(
                (contract.contract_number for contract in emp.contracts),
                default=0,
            ) + 1
            contract = models.EmploymentContract(
                employee_id=emp.id,
                contract_number=contract_number,
                start_date=payload.contract_start_date,
                end_date=payload.contract_end_date,
                status="ACTIVE",
                keterangan=(
                    f"Perubahan status PHL ke PKWT. {payload.contract_keterangan.strip()}"
                    if payload.contract_keterangan and payload.contract_keterangan.strip()
                    else "Perubahan status PHL ke PKWT."
                ),
            )
            db.add(contract)
            db.add(models.EmployeeEmploymentTransition(
                employee=emp,
                employment_contract=contract,
                phl_start_date=phl_period_start,
                phl_end_date=phl_period_end,
                transition_date=payload.contract_start_date,
            ))
        elif payload.contract_start_date and payload.contract_end_date:
            if payload.contract_end_date <= payload.contract_start_date:
                raise HTTPException(
                    status_code=400,
                    detail="Tanggal selesai kontrak (contract_end_date) harus setelah tanggal mulai kontrak."
                )
            active_c = next((c for c in emp.contracts if c.status == "ACTIVE"), None)
            if active_c:
                active_c.start_date = payload.contract_start_date
                active_c.end_date = payload.contract_end_date
                if payload.contract_keterangan:
                    active_c.keterangan = payload.contract_keterangan
            else:
                next_num = len(emp.contracts) + 1
                new_c = models.EmploymentContract(
                    employee_id=emp.id,
                    contract_number=next_num,
                    start_date=payload.contract_start_date,
                    end_date=payload.contract_end_date,
                    status="ACTIVE",
                    keterangan=payload.contract_keterangan or f"Kontrak PKWT #{next_num}",
                )
                db.add(new_c)

    emp_dict = payload.dict(
        exclude={
            "contract_start_date",
            "contract_end_date",
            "contract_keterangan",
            "phl_period_start_date",
        },
        exclude_unset=True,
    )
    if "employment_status" in emp_dict:
        emp_dict["employment_status"] = _normalize_employment_status(emp_dict["employment_status"])
    for field, value in emp_dict.items():
        setattr(emp, field, value)

    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        if is_phl_to_pkwt:
            raise HTTPException(
                status_code=409,
                detail="Transisi PHL ke PKWT tidak tersimpan karena konflik data; tidak ada perubahan yang diterapkan.",
            ) from exc
        raise
    db.refresh(emp)
    _enrich_employee_data(emp)
    return emp


@router.post("/{employee_id}/contracts", response_model=schemas.EmploymentContractOut)
def add_contract_renewal(
    employee_id: int,
    payload: schemas.EmploymentContractCreate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_hr_master),
):
    """
    Menambahkan perpanjangan kontrak baru (Kontrak ke-2, ke-3, dst.) untuk karyawan PKWT.
    Kontrak lama akan ditandai RENEWED, dan kontrak baru dibuat dengan status ACTIVE.
    """
    emp = db.query(models.Employee).filter(models.Employee.id == employee_id).first()
    if not emp:
        raise HTTPException(status_code=404, detail="Karyawan tidak ditemukan.")

    if payload.end_date <= payload.start_date:
        raise HTTPException(
            status_code=400,
            detail="Tanggal selesai kontrak (end_date) harus setelah tanggal mulai kontrak (start_date)."
        )

    # Tandai kontrak aktif sebelumnya sebagai RENEWED
    for c in emp.contracts:
        if c.status == "ACTIVE":
            c.status = "RENEWED"

    next_number = max([c.contract_number for c in emp.contracts], default=0) + 1

    # Pastikan status kepegawaian PKWT
    emp.employment_status = "PKWT"

    new_contract = models.EmploymentContract(
        employee_id=emp.id,
        contract_number=next_number,
        start_date=payload.start_date,
        end_date=payload.end_date,
        status="ACTIVE",
        keterangan=payload.keterangan or f"Perpanjangan Kontrak ke-{next_number}",
    )
    db.add(new_contract)
    db.commit()
    db.refresh(new_contract)
    return new_contract


@router.delete("/{employee_id}")
def delete_employee(
    employee_id: int,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_hr_master),
):
    emp = db.query(models.Employee).filter(models.Employee.id == employee_id).first()
    if not emp:
        raise HTTPException(status_code=404, detail="Karyawan tidak ditemukan.")
    db.delete(emp)
    db.commit()
    return {"status": "sukses"}


@router.post("/import")
async def import_employees(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.require_hr_master),
):
    from ..employee_imports import normalize_text, parse_employee_workbook

    try:
        rows, errors = parse_employee_workbook(await file.read(), file.filename or "")
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    profiles = {
        profile.code.strip().casefold(): profile.code
        for profile in db.query(models.Profile).all()
    }
    employees_by_code = {
        employee.employee_code: employee
        for employee in db.query(models.Employee).all()
        if employee.employee_code
    }
    employees = db.query(models.Employee).all()
    inserted = 0
    updated = 0
    skipped = 0
    contract_required = []
    seen_niks = {
        employee.nik: employee.id
        for employee in employees
        if employee.nik
    }

    for row in rows:
        code = row["employee_code"]
        profile_code = profiles.get(row["profile_code"].casefold())
        if not profile_code:
            errors.append({
                "row": row["source_row"],
                "employee_code": code,
                "reason": f"Kode Profil '{row['profile_code']}' tidak ditemukan.",
            })
            continue

        employee = employees_by_code.get(code)
        if employee is None and not row["join_date"]:
            errors.append({
                "row": row["source_row"],
                "employee_code": code,
                "reason": "Tanggal Masuk wajib diisi untuk karyawan baru.",
            })
            continue

        if row["nik"] and row["nik"] in seen_niks and (employee is None or seen_niks[row["nik"]] != employee.id):
            errors.append({
                "row": row["source_row"],
                "employee_code": code,
                "reason": f"NIK sudah digunakan oleh karyawan lain ({row['nik']}).",
            })
            continue

        if employee is not None and normalize_text(employee.nama) != normalize_text(row["nama"]):
            conflict = _name_reference_conflict(db, employee, row["nama"])
            if conflict:
                errors.append({
                    "row": row["source_row"],
                    "employee_code": code,
                    "reason": conflict,
                })
                continue
            _rename_employee_name_references(db, employee.nama, row["nama"])
            employee.nama = row["nama"]

        if employee is None:
            employee = models.Employee(
                employee_code=code,
                nama=row["nama"],
                profile_code=profile_code,
                cabang=row["cabang"],
                nik=row["nik"],
                id_mesin=row["id_mesin"],
                jabatan=row["jabatan"],
                employment_status=row["employment_status"],
                join_date=row["join_date"],
                active=row["active"],
                profile_needs_review=False,
            )
            db.add(employee)
            db.flush()
            employees_by_code[code] = employee
            inserted += 1
        else:
            employee.profile_code = profile_code
            employee.cabang = row["cabang"]
            employee.employment_status = row["employment_status"]
            if row["nik"]:
                employee.nik = row["nik"]
            if row["join_date"]:
                employee.join_date = row["join_date"]
            if row["id_mesin"]:
                employee.id_mesin = row["id_mesin"]
            if row["jabatan"]:
                employee.jabatan = row["jabatan"]
            if row["active_provided"]:
                employee.active = row["active"]
            updated += 1

        if row["nik"]:
            seen_niks[row["nik"]] = employee.id
        active_contract = next((contract for contract in employee.contracts if contract.status == "ACTIVE"), None)
        if row["employment_status"] == "PKWT" and active_contract is None:
            employee.active = False
            contract_required.append(code)

    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="Impor dibatalkan karena Kode Karyawan atau NIK duplikat.") from exc
    except Exception:
        db.rollback()
        raise

    skipped = len(errors)
    return {
        "status": "sukses",
        "total": len(rows) + len(errors),
        "inserted": inserted,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
        "contract_required": contract_required,
        "pesan": f"Berhasil memproses {len(rows) + len(errors)} baris: {inserted} baru, {updated} diperbarui, {skipped} perlu diperiksa.",
    }
