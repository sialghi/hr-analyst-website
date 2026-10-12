# -*- coding: utf-8 -*-
import os
import shutil
import tempfile
import uuid
import datetime
import math

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from .. import models, auth, schemas
from ..database import get_db
from ..pipeline.service import jalankan_pipeline_db_json
from ..leave_logic import get_ledger_balance, record_leave_used
from ..shared_store import _file_store

router = APIRouter(tags=["process"])

ALLOWED_EXT = (".xlsx", ".xls", ".csv")


def _simpan_upload(f: UploadFile, folder: str, nama_file: str = None) -> str:
    ekstensi = os.path.splitext(f.filename or "")[1].lower()
    if ekstensi not in ALLOWED_EXT:
        raise HTTPException(status_code=400, detail=f"File '{f.filename}' harus berformat .xlsx/.xls/.csv")
    path_tujuan = os.path.join(folder, nama_file or f.filename)
    with open(path_tujuan, "wb") as out:
        shutil.copyfileobj(f.file, out)
    return path_tujuan


def _create_upload_batch(db: Session, filenames, uploaded_by=None):
    batch = models.AttendanceUploadBatch(
        filename=", ".join(name or "absensi" for name in filenames),
        uploaded_by=uploaded_by,
        status="PROCESSING",
    )
    db.add(batch)
    db.commit()
    db.refresh(batch)
    return batch


def _finish_upload_batch(db: Session, batch, result_data=None, row_count=None, status="COMPLETED"):
    if result_data:
        period = result_data.get("periode") or {}
        start_date = period.get("start_date")
        end_date = period.get("end_date")
        batch.period_start = datetime.date.fromisoformat(start_date) if start_date else None
        batch.period_end = datetime.date.fromisoformat(end_date) if end_date else None
    if row_count is not None:
        batch.row_count = int(row_count)
    batch.status = status
    db.commit()


# ===========================================================================
# 1) ENDPOINT MANUAL DARI WEB — mengembalikan JSON lengkap + file_id download
# ===========================================================================

@router.post("/proses-absensi-manual")
async def proses_absensi_manual(
    files: list[UploadFile] = File(..., description="Satu atau lebih file absensi mentah"),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.get_current_user),
):
    """Upload manual lewat web: master karyawan & tanggal merah otomatis dari database.
    Mengembalikan JSON terstruktur lengkap dengan data 14 sheet + file_id untuk download Excel.
    """
    if not files:
        raise HTTPException(status_code=400, detail="Tidak ada file absensi yang dikirim.")

    folder_kerja = tempfile.mkdtemp(prefix="hr_pipeline_")
    path_input_list = [_simpan_upload(f, folder_kerja) for f in files]
    batch = _create_upload_batch(db, [f.filename for f in files], uploaded_by=current_user.nama)

    nama_dasar = os.path.splitext(files[0].filename or "")[0] if files else "Absensi"
    nama_download = f"Report Hasil_{nama_dasar}.xlsx"
    file_id = uuid.uuid4().hex
    path_output = os.path.join(folder_kerja, f"Report_Hasil_{file_id[:8]}.xlsx")

    try:
        result_data, _ = jalankan_pipeline_db_json(
            daftar_input=path_input_list,
            path_output=path_output,
            db=db,
            upload_batch_id=batch.id,
            verbose=False,
        )
        _finish_upload_batch(
            db,
            batch,
            result_data=result_data,
        )
    except Exception as e:
        _finish_upload_batch(db, batch, status="FAILED")
        raise HTTPException(status_code=422, detail=f"Gagal memproses data: {e}")

    # Simpan referensi file untuk download nanti (juga simpan result_data untuk chatbot)
    _file_store[file_id] = {"path": path_output, "filename": nama_download, "result_data": result_data}

    # Tambahkan file_id & filename ke response
    result_data["file_id"] = file_id
    result_data["filename"] = nama_download
    result_data["upload_batch_id"] = batch.id

    return JSONResponse(content=result_data)


@router.get("/proses/download/{file_id}")
async def download_hasil_proses(
    file_id: str,
    _: models.User = Depends(auth.get_current_user),
):
    """Download file Excel hasil proses berdasarkan file_id."""
    entry = _file_store.get(file_id)
    if not entry or not os.path.exists(entry["path"]):
        raise HTTPException(status_code=404, detail="File tidak ditemukan atau sudah kedaluwarsa.")
    return FileResponse(
        entry["path"],
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=entry["filename"],
    )


@router.patch("/proses/{file_id}/adjustments")
def update_processed_adjustment(
    file_id: str,
    payload: schemas.EmployeeAdjustmentUpdate,
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_user),
):
    """Save an employee adjustment and synchronize the active report and Excel file."""
    entry = _file_store.get(file_id)
    if not entry or not os.path.isfile(entry.get("path", "")):
        raise HTTPException(status_code=404, detail="File proses tidak ditemukan atau sudah kedaluwarsa.")

    nama = payload.nama.strip()
    if not nama:
        raise HTTPException(status_code=400, detail="Nama karyawan wajib diisi.")
    if any(
        value is not None and not math.isfinite(value)
        for value in (payload.bonus_lain, payload.potongan_lain)
    ):
        raise HTTPException(status_code=400, detail="Nominal bonus dan potongan harus berupa angka valid.")

    result_data = entry.get("result_data") or {}
    summary_sheet = next(
        (sheet for sheet in result_data.get("sheets", []) if sheet.get("key") == "summary"),
        None,
    )
    summary_rows = summary_sheet.get("data", {}).get("rows", []) if summary_sheet else []
    matching_rows = [
        row for row in summary_rows
        if str(row.get("Nama") or "").strip().casefold() == nama.casefold()
    ]
    if not matching_rows:
        raise HTTPException(status_code=404, detail="Karyawan tidak ditemukan pada ringkasan file proses ini.")

    adjustment = (
        db.query(models.EmployeeAdjustment)
        .filter(models.EmployeeAdjustment.nama == nama)
        .first()
    )
    stored_bonus = float(adjustment.bonus_lain or 0) if adjustment else float(
        matching_rows[0].get("Total Bonus Lain-lain (Rp)") or 0
    )
    stored_deduction = float(adjustment.potongan_lain or 0) if adjustment else float(
        matching_rows[0].get("Total Potongan Lain-lain (Rp)") or 0
    )
    new_bonus = payload.bonus_lain if payload.bonus_lain is not None else stored_bonus
    new_deduction = payload.potongan_lain if payload.potongan_lain is not None else stored_deduction

    report_path = entry["path"]
    try:
        workbook = load_workbook(report_path)
    except Exception as exc:
        raise HTTPException(status_code=500, detail="Workbook hasil proses tidak dapat dibuka untuk diperbarui.") from exc

    sheet_name = "2. Summary Overview"
    if sheet_name not in workbook.sheetnames:
        raise HTTPException(status_code=500, detail="Sheet ringkasan tidak ditemukan pada workbook hasil proses.")
    worksheet = workbook[sheet_name]
    column_indexes = {
        str(cell.value).strip(): cell.column
        for cell in worksheet[4]
        if cell.value is not None
    }
    required_columns = (
        "Nama",
        "Total Bonus Lain-lain (Rp)",
        "Total Potongan Lain-lain (Rp)",
        "Total Bersih (Rp)",
    )
    if any(column not in column_indexes for column in required_columns):
        raise HTTPException(status_code=500, detail="Kolom penyesuaian/Total Bersih tidak lengkap pada workbook.")

    matching_excel_rows = [
        row_number
        for row_number in range(5, worksheet.max_row + 1)
        if str(worksheet.cell(row_number, column_indexes["Nama"]).value or "").strip().casefold() == nama.casefold()
    ]
    if len(matching_excel_rows) != len(matching_rows):
        raise HTTPException(status_code=409, detail="Baris ringkasan aplikasi dan workbook tidak cocok.")

    temporary_path = None
    try:
        for row_number in matching_excel_rows:
            worksheet.cell(row_number, column_indexes["Total Bonus Lain-lain (Rp)"]).value = new_bonus
            worksheet.cell(row_number, column_indexes["Total Potongan Lain-lain (Rp)"]).value = new_deduction
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True

        with tempfile.NamedTemporaryFile(
            suffix=".xlsx",
            prefix="hr_adjustment_",
            dir=os.path.dirname(report_path),
            delete=False,
        ) as temporary_file:
            temporary_path = temporary_file.name
        workbook.save(temporary_path)

        if adjustment is None:
            adjustment = models.EmployeeAdjustment(
                nama=nama,
                bonus_lain=new_bonus,
                potongan_lain=new_deduction,
            )
            db.add(adjustment)
        else:
            adjustment.bonus_lain = new_bonus
            adjustment.potongan_lain = new_deduction
        db.commit()
    except Exception as exc:
        db.rollback()
        if temporary_path and os.path.exists(temporary_path):
            os.remove(temporary_path)
        raise HTTPException(status_code=500, detail="Gagal menyimpan penyesuaian karyawan.") from exc

    try:
        os.replace(temporary_path, report_path)
        temporary_path = None
    except OSError as exc:
        if temporary_path and os.path.exists(temporary_path):
            os.remove(temporary_path)
        raise HTTPException(
            status_code=500,
            detail="Penyesuaian tersimpan, tetapi workbook hasil proses gagal diperbarui. Simpan ulang untuk mencoba sinkronisasi lagi.",
        ) from exc

    for row in matching_rows:
        report_old_bonus = float(row.get("Total Bonus Lain-lain (Rp)") or 0)
        report_old_deduction = float(row.get("Total Potongan Lain-lain (Rp)") or 0)
        row["Total Bonus Lain-lain (Rp)"] = new_bonus
        row["Total Potongan Lain-lain (Rp)"] = new_deduction
        row["Total Bersih (Rp)"] = (
            float(row.get("Total Bersih (Rp)") or 0)
            + (new_bonus - report_old_bonus)
            - (new_deduction - report_old_deduction)
        )
    result_data.setdefault("kpis", {})["total_bersih"] = sum(
        float(row.get("Total Bersih (Rp)") or 0) for row in summary_rows
    )

    return {
        "status": "sukses",
        "nama": nama,
        "bonus_lain": new_bonus,
        "potongan_lain": new_deduction,
        "updated_rows": len(matching_rows),
        "total_bersih": result_data["kpis"]["total_bersih"],
    }


@router.get("/attendance/upload-batches")
def list_attendance_upload_batches(
    limit: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_user),
):
    batches = (
        db.query(models.AttendanceUploadBatch)
        .order_by(models.AttendanceUploadBatch.created_at.desc())
        .limit(limit)
        .all()
    )
    return [{
        "id": batch.id,
        "filename": batch.filename,
        "period_start": batch.period_start,
        "period_end": batch.period_end,
        "row_count": batch.row_count,
        "status": batch.status,
        "uploaded_by": batch.uploaded_by,
        "created_at": batch.created_at,
    } for batch in batches]


@router.get("/attendance/daily")
def list_attendance_daily(
    employee_id: int | None = Query(None),
    start_date: datetime.date | None = Query(None),
    end_date: datetime.date | None = Query(None),
    limit: int = Query(500, ge=1, le=5000),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_user),
):
    query = db.query(models.AttendanceDaily).order_by(
        models.AttendanceDaily.attendance_date.desc(),
        models.AttendanceDaily.employee_id,
    )
    if employee_id is not None:
        query = query.filter(models.AttendanceDaily.employee_id == employee_id)
    if start_date is not None:
        query = query.filter(models.AttendanceDaily.attendance_date >= start_date)
    if end_date is not None:
        query = query.filter(models.AttendanceDaily.attendance_date <= end_date)

    rows = query.limit(limit).all()
    return [{
        "id": row.id,
        "employee_id": row.employee_id,
        "tanggal": row.attendance_date,
        "status": row.status,
        "cabang": row.cabang,
        "profile": row.profile,
        "jam_masuk": row.jam_masuk,
        "jam_keluar": row.jam_keluar,
        "kategori_hari": row.kategori_hari,
        "catatan": row.catatan,
        "upload_batch_id": row.upload_batch_id,
        "updated_at": row.updated_at,
    } for row in rows]


@router.get("/attendance/reviews")
def list_attendance_reviews(
    status: str = Query("PENDING", pattern="^(PENDING|RESOLVED|ALL)$"),
    upload_batch_id: int | None = Query(None, ge=1),
    limit: int = Query(500, ge=1, le=5000),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_user),
):
    query = db.query(models.AttendanceReview).order_by(
        models.AttendanceReview.attendance_date.desc(),
        models.AttendanceReview.id.desc(),
    )
    if status != "ALL":
        query = query.filter(models.AttendanceReview.status == status)
    if upload_batch_id is not None:
        query = query.filter(models.AttendanceReview.upload_batch_id == upload_batch_id)
    rows = query.limit(limit).all()
    from ..config import ANNUAL_LEAVE_QUOTA

    return [{
        "id": row.id,
        "employee_id": row.employee_id,
        "nama": row.employee.nama if row.employee else None,
        "employment_status": row.employee.employment_status if row.employee else None,
        "tanggal": row.attendance_date,
        "status": row.status,
        "decision": row.decision,
        "note": row.note,
        "upload_batch_id": row.upload_batch_id,
        "created_at": row.created_at,
        "resolved_at": row.resolved_at,
        "resolved_by": row.resolved_by,
        "sisa_cuti_tahunan": max(
            0.0,
            get_ledger_balance(db, row.employee_id, row.attendance_date.year),
        ),
        "kuota_cuti_tahunan": ANNUAL_LEAVE_QUOTA,
    } for row in rows]


@router.patch("/attendance/reviews/{review_id}")
def resolve_attendance_review(
    review_id: int,
    payload: dict,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(auth.require_hr_master),
):
    review = db.query(models.AttendanceReview).filter(
        models.AttendanceReview.id == review_id
    ).first()
    if not review:
        raise HTTPException(status_code=404, detail="Review ketidakhadiran tidak ditemukan.")
    if review.status != "PENDING":
        raise HTTPException(status_code=400, detail="Review ini sudah diselesaikan.")

    decision = str(payload.get("decision") or "").strip().upper()
    note = str(payload.get("note") or "").strip() or None
    employee = review.employee
    if employee and str(employee.employment_status or "").strip().upper() == "PHL" and decision != "UNPAID_LEAVE":
        raise HTTPException(
            status_code=400,
            detail="Untuk karyawan PHL, review absensi hanya dapat diselesaikan sebagai Unpaid Leave.",
        )
    allowed = {"ANNUAL_LEAVE", "UNPAID_LEAVE", "SICK_OR_PERMISSION", "CORRECTED", "IGNORED"}
    if decision not in allowed:
        raise HTTPException(status_code=400, detail="Keputusan tidak valid.")

    leave_request_id = None
    if decision == "ANNUAL_LEAVE":
        balance = max(0.0, get_ledger_balance(db, employee.id, review.attendance_date.year))
        if balance < 1:
            raise HTTPException(status_code=400, detail="Saldo cuti tahunan tidak mencukupi. Pilih Unpaid Leave atau koreksi.")
        leave = models.LeaveRequest(
            nama=employee.nama,
            kategori="CUTI_TAHUNAN",
            tanggal_mulai=review.attendance_date,
            tanggal_selesai=review.attendance_date,
            jumlah_hari=1.0,
            unpaid_leave_days=0.0,
            alasan="Ditetapkan dari review absensi tanpa scan.",
            catatan_hr=note,
            status="APPROVED",
            approved_by=f"HR Master ({current_user.nama})",
            approved_at=datetime.datetime.utcnow(),
        )
        db.add(leave)
        db.flush()
        record_leave_used(
            db,
            employee.id,
            review.attendance_date.year,
            1.0,
            leave.id,
            f"CUTI_TAHUNAN otomatis dari review {review.attendance_date}",
        )
        leave_request_id = leave.id
    elif decision in {"UNPAID_LEAVE", "SICK_OR_PERMISSION"}:
        leave = models.LeaveRequest(
            nama=employee.nama,
            kategori="UNPAID_LEAVE" if decision == "UNPAID_LEAVE" else "SAKIT",
            tanggal_mulai=review.attendance_date,
            tanggal_selesai=review.attendance_date,
            jumlah_hari=1.0,
            unpaid_leave_days=1.0 if decision == "UNPAID_LEAVE" else 0.0,
            alasan="Ditetapkan dari review absensi tanpa scan.",
            catatan_hr=note,
            status="APPROVED",
            approved_by=f"HR Master ({current_user.nama})",
            approved_at=datetime.datetime.utcnow(),
        )
        db.add(leave)
        db.flush()
        leave_request_id = leave.id

    review.status = "RESOLVED"
    review.decision = decision
    review.note = note
    review.leave_request_id = leave_request_id
    review.resolved_at = datetime.datetime.utcnow()
    review.resolved_by = current_user.nama
    daily = db.query(models.AttendanceDaily).filter(
        models.AttendanceDaily.employee_id == employee.id,
        models.AttendanceDaily.attendance_date == review.attendance_date,
    ).first()
    if daily:
        daily.status = decision
        daily.catatan = note or f"Diselesaikan HR sebagai {decision}."
    db.commit()
    db.refresh(review)
    return {
        "id": review.id,
        "status": review.status,
        "decision": review.decision,
        "leave_request_id": review.leave_request_id,
    }


@router.get("/proses/template")
def download_attendance_template(
    format: str = Query("xlsx", pattern="^(xlsx|csv)$"),
    db: Session = Depends(get_db),
    _: models.User = Depends(auth.get_current_user),
):
    """Mengunduh file template/contoh log absensi mentah (.xlsx atau .csv).
    Data contoh disesuaikan dengan karyawan aktif di master jika tersedia.
    """
    import io
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    # Ambil sampel karyawan aktif dari database
    emps = db.query(models.Employee).filter(models.Employee.active == True).limit(5).all()
    if emps:
        sample_emps = [(e.nama, e.cabang or "OFFICE") for e in emps]
    else:
        sample_emps = [
            ("Budi Hartono", "BLOK C"),
            ("Andi Santanu", "BLOK C"),
            ("Dian Maulana", "BLOK C"),
            ("Ahmad Fauzi", "BANDUNG"),
        ]

    # Buat baris log contoh realistis (2 hari scan: normal, lembur, telat)
    rows = []
    for idx, (nama, cabang) in enumerate(sample_emps):
        if idx == 0:
            # Jam kerja normal
            rows.append([cabang, nama, "1", "02/03/2026 07:54:12", "C/Masuk", "1", "", "Finger"])
            rows.append([cabang, nama, "1", "02/03/2026 16:05:30", "C/Keluar", "1", "", "Finger"])
            rows.append([cabang, nama, "1", "03/03/2026 07:50:45", "C/Masuk", "1", "", "Finger"])
            rows.append([cabang, nama, "1", "03/03/2026 16:12:18", "C/Keluar", "1", "", "Finger"])
        elif idx == 1:
            # Lembur (pulang malam)
            rows.append([cabang, nama, "1", "02/03/2026 07:58:00", "C/Masuk", "1", "", "Finger"])
            rows.append([cabang, nama, "1", "02/03/2026 18:35:10", "C/Keluar", "1", "", "Finger"])
            rows.append([cabang, nama, "1", "03/03/2026 07:48:22", "C/Masuk", "1", "", "Finger"])
            rows.append([cabang, nama, "1", "03/03/2026 16:02:40", "C/Keluar", "1", "", "Finger"])
        elif idx == 2:
            # Telat sedikit (datang lewat 08:00)
            rows.append([cabang, nama, "1", "02/03/2026 08:22:15", "C/Masuk", "1", "", "Finger"])
            rows.append([cabang, nama, "1", "02/03/2026 16:01:05", "C/Keluar", "1", "", "Finger"])
            rows.append([cabang, nama, "1", "03/03/2026 07:55:00", "C/Masuk", "1", "", "Finger"])
            rows.append([cabang, nama, "1", "03/03/2026 16:08:44", "C/Keluar", "1", "", "Finger"])
        else:
            rows.append([cabang, nama, "1", "02/03/2026 07:52:00", "C/Masuk", "1", "", "Finger"])
            rows.append([cabang, nama, "1", "02/03/2026 16:06:00", "C/Keluar", "1", "", "Finger"])
            rows.append([cabang, nama, "1", "03/03/2026 07:56:00", "C/Masuk", "1", "", "Finger"])
            rows.append([cabang, nama, "1", "03/03/2026 16:04:00", "C/Keluar", "1", "", "Finger"])

    columns = [
        "Cabang",
        "Nama",
        "ID_Mesin",
        "Timestamp_Raw",
        "Tipe_Asli_Raw",
        "Kode_Gerbang",
        "Kolom_Kosong",
        "Metode",
    ]

    df = pd.DataFrame(rows, columns=columns)

    if format == "csv":
        out_csv = io.StringIO()
        df.to_csv(out_csv, index=False)
        return Response(
            content=out_csv.getvalue().encode("utf-8-sig"),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": 'attachment; filename="template_absensi_mentah.csv"'},
        )

    # Format Excel (.xlsx) dengan 2 Sheet: Data Absensi & Panduan Kolom
    wb = Workbook()
    ws_data = wb.active
    ws_data.title = "Log Absensi"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11, color="0F172A")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    ws_data.append(columns)
    for col_num in range(1, len(columns) + 1):
        cell = ws_data.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

    for r_idx, row_data in enumerate(rows, start=2):
        ws_data.append(row_data)
        use_zebra = (r_idx % 2 == 1)
        for c_idx in range(1, len(row_data) + 1):
            cell = ws_data.cell(row=r_idx, column=c_idx)
            cell.font = data_font
            cell.border = thin_border
            if use_zebra:
                cell.fill = zebra_fill
            if c_idx in (3, 6):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 4:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws_data.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_data.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Sheet 2: Panduan & Penjelasan Format Kolom
    ws_guide = wb.create_sheet(title="Panduan Kolom")
    ws_guide.views.sheetView[0].showGridLines = True

    guide_headers = ["No", "Nama Kolom", "Wajib?", "Contoh Nilai", "Keterangan & Aturan"]
    guide_rows = [
        [1, "Cabang", "Ya", "BANDUNG / BLOK C / BLOK A", "Lokasi atau cabang tempat scan absensi dilakukan."],
        [2, "Nama", "Ya", "Budi Hartono", "Nama lengkap karyawan. WAJIB sama dengan nama di menu Master Karyawan."],
        [3, "ID_Mesin", "Opsional", "1", "Nomor atau ID unit mesin absensi."],
        [4, "Timestamp_Raw", "Ya", "02/03/2026 07:54:12", "Waktu scan. Format disarankan: DD/MM/YYYY HH:MM:SS atau YYYY-MM-DD HH:MM:SS."],
        [5, "Tipe_Asli_Raw", "Ya", "C/Masuk atau C/Keluar", "Arah scan asli mesin. Sistem juga mengoreksi otomatis jika label salah/terbalik."],
        [6, "Kode_Gerbang", "Opsional", "1", "Kode pintu gerbang / terminal."],
        [7, "Kolom_Kosong", "Opsional", "", "Kolom kosong bawaan format log mesin absensi."],
        [8, "Metode", "Opsional", "Finger / Face / Password", "Metode verifikasi absensi."],
    ]

    title_cell = ws_guide.cell(row=1, column=1, value="PANDUAN FORMAT LOG ABSENSI MENTAH (8 KOLOM)")
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1E293B")

    sub_cell = ws_guide.cell(
        row=2, column=1,
        value="File log mentah dapat berupa .xlsx, .xls, maupun .csv. Sistem otomatis mendeteksi file dengan atau tanpa baris header."
    )
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="64748B")

    ws_guide.append([])
    ws_guide.append(guide_headers)
    for col_num in range(1, len(guide_headers) + 1):
        cell = ws_guide.cell(row=4, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, grow in enumerate(guide_rows, start=5):
        ws_guide.append(grow)
        use_zebra = (r_idx % 2 == 1)
        for c_idx in range(1, len(grow) + 1):
            cell = ws_guide.cell(row=r_idx, column=c_idx)
            cell.font = data_font
            cell.border = thin_border
            if use_zebra:
                cell.fill = zebra_fill
            if c_idx in (1, 3):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws_guide.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_guide.column_dimensions[col_letter].width = max(min(max_len + 4, 60), 10)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="template_absensi_mentah.xlsx"'},
    )
