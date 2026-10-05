# -*- coding: utf-8 -*-
"""
Script Import Data Kehadiran 2026 dari 'Rekap kehadiran 2026.xlsx' ke database HR Analyst.

Aturan Import:
1. 'C'  -> Kategori: CUTI_TAHUNAN, Status: APPROVED, Jumlah Hari: 1.0
2. 'CS' -> Kategori: CUTI_SETENGAH_HARI, Status: APPROVED, Jumlah Hari: 0.5
3. Data bersifat idempotent (tidak membuat duplikasi jika dijalankan ulang).
"""

import os
import glob
import datetime
import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

# Set absolute path database SQLite backend
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "hr_app.db")

from app.database import Base
from app import models

engine = create_engine(f"sqlite:///{DB_PATH}", connect_args={"check_same_thread": False})
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)


def run_import():
    db = SessionLocal()
    try:
        # Cari file Rekap kehadiran 2026.xlsx
        project_root = os.path.abspath(os.path.join(BASE_DIR, ".."))
        matches = glob.glob(os.path.join(project_root, "**", "Rekap kehadiran 2026*.xlsx"), recursive=True)
        if not matches:
            print("[ERROR] File 'Rekap kehadiran 2026.xlsx' tidak ditemukan!")
            return

        excel_path = matches[0]
        print(f"[INFO] Membaca file: {excel_path}")

        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        sheet_name = "Absensi" if "Absensi" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]

        # Read row 3 for dates mapping
        date_mapping = {}  # col_idx -> date_object
        row3 = None
        
        c_count = 0
        cs_count = 0
        skipped_count = 0

        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if row_idx == 3:
                row3 = row
                for c_idx in range(3, len(row)):
                    val = row[c_idx]
                    if isinstance(val, datetime.datetime):
                        date_mapping[c_idx] = val.date()
                    elif isinstance(val, datetime.date):
                        date_mapping[c_idx] = val
                print(f"[INFO] Ditemukan {len(date_mapping)} kolom tanggal di header Excel.")
                continue

            if row_idx < 5:
                continue

            # Row 5 onwards: Data Karyawan
            no_val = row[0]
            nama = row[2]
            if not nama or str(nama).strip() == "":
                continue

            nama_clean = str(nama).strip()

            for c_idx, tgl in date_mapping.items():
                if c_idx >= len(row):
                    continue
                cell_val = str(row[c_idx]).strip().upper() if row[c_idx] is not None else ""

                kategori = None
                jumlah_hari = None
                alasan = None

                if cell_val == "C":
                    kategori = "CUTI_TAHUNAN"
                    jumlah_hari = 1.0
                    alasan = "Cuti Tahunan (Import Rekap Kehadiran 2026)"
                elif cell_val == "CS":
                    kategori = "CUTI_SETENGAH_HARI"
                    jumlah_hari = 0.5
                    alasan = "Cuti Setengah Hari (Import Rekap Kehadiran 2026)"

                if kategori:
                    # Cek keberadaan record agar idempotent
                    existing = (
                        db.query(models.LeaveRequest)
                        .filter(
                            models.LeaveRequest.nama.ilike(nama_clean),
                            models.LeaveRequest.tanggal_mulai == tgl,
                            models.LeaveRequest.kategori == kategori,
                        )
                        .first()
                    )

                    if not existing:
                        new_leave = models.LeaveRequest(
                            nama=nama_clean,
                            kategori=kategori,
                            tanggal_mulai=tgl,
                            tanggal_selesai=tgl,
                            jumlah_hari=jumlah_hari,
                            alasan=alasan,
                            status="APPROVED",
                            catatan_hr="Diimpor otomatis dari Rekap Kehadiran 2026.xlsx",
                            approved_by="System Import 2026",
                            approved_at=datetime.datetime.utcnow(),
                        )
                        db.add(new_leave)

                        if cell_val == "C":
                            c_count += 1
                        else:
                            cs_count += 1
                    else:
                        skipped_count += 1

        db.commit()
        print("\n=== HASIL IMPORT REKAP KEHADIRAN 2026 ===")
        print(f"[SUCCESS] Record 'C' (1.0 hari) baru diimport   : {c_count}")
        print(f"[SUCCESS] Record 'CS' (0.5 hari) baru diimport  : {cs_count}")
        print(f"[INFO]    Record sudah ada (skipped)          : {skipped_count}")

        # Tampilkan Rekap Saldo Cuti 2026 untuk karyawan yang melebihi kuota 12 hari
        print("\n=== REKAP KARYAWAN MELEBIHI KUOTA CUTI 2026 (> 12 HARI) ===")
        emp_names = set(r[0] for r in db.query(models.LeaveRequest.nama).all() if r[0])
        
        exceeded_list = []
        for name in sorted(emp_names):
            leaves = (
                db.query(models.LeaveRequest)
                .filter(
                    models.LeaveRequest.nama.ilike(name),
                    models.LeaveRequest.kategori.in_(["CUTI_TAHUNAN", "CUTI_SETENGAH_HARI"]),
                    models.LeaveRequest.status.in_(["APPROVED", "PENDING"]),
                    models.LeaveRequest.tanggal_mulai >= datetime.date(2026, 1, 1),
                    models.LeaveRequest.tanggal_mulai <= datetime.date(2026, 12, 31),
                )
                .all()
            )

            total_used = sum(l.jumlah_hari if l.jumlah_hari is not None else (1.0 if l.kategori == "CUTI_TAHUNAN" else 0.5) for l in leaves)
            if total_used > 12:
                exceeded_list.append((name, total_used, round(total_used - 12, 1)))

        exceeded_list.sort(key=lambda x: x[1], reverse=True)
        print(f"Total karyawan melebihi kuota: {len(exceeded_list)} orang\n")
        print(f"{'Nama Karyawan':<35} | {'Total Dipakai':<15} | {'Kelebihan (Hari)':<15}")
        print("-" * 70)
        for name, used, excess in exceeded_list:
            print(f"{name:<35} | {used:<15.1f} | +{excess:<14.1f}")

    finally:
        db.close()


if __name__ == "__main__":
    run_import()
